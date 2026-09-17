import pytest
pytest.skip("legacy docx_template/pdf_template removed; replaced by template_compliance", allow_module_level=True)

"""DOCX/PDF SOP structure extraction integration tests (extractor-level, no Neo4j)."""

from __future__ import annotations

from pathlib import Path

import pytest


@pytest.fixture()
def sop_docx(tmp_path: Path) -> Path:
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")

    doc = Document()
    doc.add_heading("SOP-001 Cleaning Procedure", level=0)
    doc.add_heading("1 Purpose", level=1)
    doc.add_paragraph("Describe the cleaning procedure.")
    doc.add_heading("4 Responsibilities", level=1)
    doc.add_paragraph("QA: Approve the procedure")
    doc.add_paragraph("QC: Verify cleanliness")
    doc.add_paragraph("Production: Execute cleaning")
    doc.add_heading("5 Procedure", level=1)
    doc.add_paragraph("5.1 Switch ON machine")
    doc.add_paragraph("5.2 Verify pressure")
    doc.add_paragraph("5.3 Record observations")
    doc.add_paragraph("1. Gather materials")
    doc.add_paragraph("2. Start cycle")
    doc.add_paragraph("A. Inspect vessel")
    doc.add_paragraph("B. Close lid")
    doc.add_heading("6 Approval", level=1)
    doc.add_paragraph("Prepared By: Alice")
    doc.add_paragraph("Reviewed By: Bob")
    doc.add_paragraph("Approved By: Carol")
    doc.add_paragraph("Approved Date: 2024-05-01")
    doc.add_heading("7 Revision History", level=1)
    table = doc.add_table(rows=3, cols=4)
    headers = ["Version", "Date", "Description", "Author"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    table.rows[1].cells[0].text = "1.0"
    table.rows[1].cells[1].text = "2023-01-01"
    table.rows[1].cells[2].text = "Initial release"
    table.rows[1].cells[3].text = "QA"
    table.rows[2].cells[0].text = "1.1"
    table.rows[2].cells[1].text = "2024-05-01"
    table.rows[2].cells[2].text = "Added steps"
    table.rows[2].cells[3].text = "QC"
    doc.add_paragraph("Figure 1: Equipment Layout")

    out = tmp_path / "sop_sample.docx"
    doc.save(str(out))
    return out


def test_docx_sop_enrichment_fields(sop_docx: Path):
    from src.features.document_processing.extractors.docx_template import extract_template

    result = extract_template(sop_docx)

    # Existing fields unchanged
    assert result["document_tree"]
    assert any(n.get("type") == "heading" for n in result["document_tree"])
    assert result["tables"]

    # New SOP fields
    assert result["procedure_steps"], "expected procedure steps"
    assert any(s["step_number"] == "5.1" for s in result["procedure_steps"])
    assert result["responsibilities"]
    assert {r["role"].upper() for r in result["responsibilities"]} >= {"QA", "QC", "PRODUCTION"}
    assert result["approvals"]
    assert result["approvals"][0]["prepared_by"] == "Alice"
    assert result["revisions"]
    assert result["revisions"][0]["version"] == "1.0"
    assert any(i.get("figure_number") == "1" for i in result["images"])


def test_docx_without_sop_structures_still_succeeds(tmp_path: Path):
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")

    doc = Document()
    doc.add_heading("Simple Doc", level=1)
    doc.add_paragraph("No procedure, roles, approvals, or revisions here.")
    path = tmp_path / "plain.docx"
    doc.save(str(path))

    from src.features.document_processing.extractors.docx_template import extract_template

    result = extract_template(path)
    assert result["document_tree"]
    assert result["procedure_steps"] == []
    assert result["approvals"] == []
    assert result["revisions"] == []


def test_pdf_template_enrichment_smoke(tmp_path: Path):
    """Build a tiny text PDF if reportlab is available; otherwise skip."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
    except ImportError:
        pytest.skip("reportlab not installed")

    path = tmp_path / "sop_sample.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    y = 750
    for line in (
        "5 Procedure",
        "5.1 Switch ON machine",
        "5.2 Verify pressure",
        "Prepared By: Alice",
        "Reviewed By: Bob",
        "Approved By: Carol",
        "Figure 1 Equipment Layout",
    ):
        c.drawString(72, y, line)
        y -= 24
    c.save()

    from src.features.document_processing.extractors.pdf_template import extract_template

    result = extract_template(path)
    assert "document_tree" in result
    assert "procedure_steps" in result
    assert "approvals" in result
    assert "images" in result
