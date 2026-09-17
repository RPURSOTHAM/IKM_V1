import json
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document

from src.features.document_processing.extractors.template_compliance import (
    compliance_to_legacy_template,
    extract_template,
)
from src.infrastructure.document_databases.neo4j_store import store_template_graph


def test_complete_template_extraction_keys(tmp_path: Path) -> None:
    docx_path = tmp_path / "test_keys.docx"
    doc = Document()
    doc.add_paragraph("1. OBJECTIVE")
    doc.add_paragraph("Test paragraph text")
    doc.save(str(docx_path))

    compliance = extract_template(docx_path)
    assert "document" in compliance
    document = compliance["document"]
    assert "metadata" in document
    assert "sections" in document

    legacy = compliance_to_legacy_template(compliance)
    for key in ("metadata", "outline", "document_tree", "tables", "source", "compliance"):
        assert key in legacy
    assert legacy["source"] == "template_compliance"


def test_formatting_metadata_persistence_mocked() -> None:
    template_data = {
        "source": "template_compliance",
        "schema_version": "1.0",
        "metadata": {"title": "Test Title"},
        "document_tree": [
            {
                "type": "heading",
                "text": "1 OBJECTIVE",
                "level": 1,
                "number": "1",
                "title": "OBJECTIVE",
                "children": [],
            }
        ],
        "outline": [],
        "tables": [],
        "content_blocks": [],
        "compliance": {"document": {"metadata": {"title": "Test Title"}}},
    }

    captured_params = {}

    def _run(cypher, **params):
        if "template_json =" in cypher or "t.template_json =" in cypher:
            captured_params.update(params)
        mock_rec = MagicMock()
        mock_rec.__getitem__ = lambda s, k: 1
        mock_result = MagicMock()
        mock_result.single.return_value = mock_rec
        return mock_result

    mock_tx = MagicMock()
    mock_tx.run.side_effect = _run
    mock_session = MagicMock()
    mock_session.__enter__ = lambda s: s
    mock_session.__exit__ = MagicMock(return_value=False)
    mock_session.execute_write = lambda fn: fn(mock_tx)
    mock_driver = MagicMock()
    mock_driver.session.return_value = mock_session
    mock_gdb = MagicMock()
    mock_gdb.driver.return_value = mock_driver

    with patch.dict("sys.modules", {"neo4j": MagicMock(GraphDatabase=mock_gdb)}):
        with patch(
            "src.infrastructure.document_databases.neo4j_store._neo4j_credentials",
            return_value=("bolt://localhost:7687", "neo4j", "password"),
        ):
            with patch(
                "src.infrastructure.document_databases.neo4j_store._require_neo4j_driver",
                return_value=mock_gdb,
            ):
                store_template_graph(
                    document_id="doc-1",
                    repository_id="repo-1",
                    template_data=template_data,
                )

    assert "template_json" in captured_params
    stored = json.loads(captured_params["template_json"])
    assert stored["source"] == "template_compliance"
    assert stored["metadata"]["title"] == "Test Title"
