"""Integration tests for PDF/DOCX document loaders."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from src.features.document_processing.loaders.component_classification import (
    COMPONENT_FOOTER,
    COMPONENT_HEADER,
    COMPONENT_PARAGRAPH,
    COMPONENT_SUBTITLE,
    COMPONENT_TABLE,
    COMPONENT_TITLE,
    block_component_type,
)
from src.features.document_processing.loaders.document_text import load_document_blocks
from src.features.document_processing.loaders.docxloader import DocxLoader
from src.features.document_processing.loaders.loader import Block
from src.features.document_processing.loaders.pdfloader import PdfLoader


def _build_sample_docx(path: Path) -> None:
    doc = Document()
    header = doc.sections[0].header
    header.paragraphs[0].text = "Document Header Line"
    doc.add_paragraph("Main Document Title", style="Title")
    doc.add_paragraph("Section Overview", style="Heading 2")
    doc.add_paragraph(
        "This is a narrative paragraph with enough words to classify as body content in the document."
    )
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Requirement"
    table.cell(1, 0).text = "Voltage"
    table.cell(1, 1).text = "230V single phase supply with UPS backup for control systems"
    table.cell(2, 0).text = "Air"
    table.cell(2, 1).text = "0 to 6 bar instrument quality compressed air"
    footer = doc.sections[0].footer
    footer.paragraphs[0].text = "Footer - Page 1"
    doc.add_page_break()
    doc.add_paragraph("Second Page Body", style="Normal")
    doc.save(str(path))


def _build_sample_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    data = [["Parameter", "Requirement"], ["Voltage", "230V"], ["Air", "0 to 6 bar"]]
    table = Table(data)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    doc = SimpleDocTemplate(str(path), pagesize=letter)
    story = [
        Paragraph("Equipment Specification", styles["Title"]),
        Paragraph("2.1 System Requirements", styles["Heading2"]),
        Paragraph("The system shall provide validated sterilization cycles.", styles["Normal"]),
        table,
    ]
    doc.build(story)


def _build_pdf_with_canvas_footer(path: Path) -> None:
    from reportlab.pdfgen import canvas

    page_width, page_height = letter
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.setFont("Helvetica", 11)
    pdf.drawString(72, page_height - 72, "Body content line")
    pdf.setFont("Helvetica", 9)
    pdf.drawCentredString(page_width / 2, 36, "Page 1 of 1")
    pdf.save()


def _line_numbers_per_page(blocks) -> dict[int, list[int]]:
    by_page: dict[int, list[int]] = {}
    for block in blocks:
        by_page.setdefault(block.page, []).append(block.line_number)
    return by_page


def test_docx_loader_extracts_structured_blocks(tmp_path: Path) -> None:
    docx_path = tmp_path / "sample.docx"
    _build_sample_docx(docx_path)
    blocks, meta = load_document_blocks(docx_path)

    assert len(blocks) >= 6
    assert meta["page_count"] >= 2
    components = {block_component_type(block) for block in blocks}
    assert COMPONENT_HEADER in components
    assert COMPONENT_TITLE in components
    assert COMPONENT_SUBTITLE in components
    assert COMPONENT_TABLE in components
    assert COMPONENT_FOOTER in components
    assert COMPONENT_PARAGRAPH in components

    page_lines = _line_numbers_per_page(blocks)
    assert page_lines[1] == sorted(page_lines[1])
    assert page_lines[1][0] == 1
    assert len(set(page_lines[1])) == len(page_lines[1])
    assert page_lines[2][0] == 1

    table_blocks = [block for block in blocks if block.block_type == "table"]
    assert table_blocks
    assert "Voltage" in table_blocks[0].text
    assert table_blocks[0].page == 1
    assert table_blocks[0].line_number > 0


def test_docx_loader_prefers_layout_pdf_blocks_for_page_lines(tmp_path: Path) -> None:
    docx_path = tmp_path / "layout.docx"
    _build_sample_docx(docx_path)
    pdf_path = tmp_path / "layout.pdf"
    pdf_path.write_bytes(b"%PDF-1.4\n")
    pdf_blocks = [
        Block(text="Page one body", page=1, line_number=1),
        Block(text="Page two body", page=2, line_number=1),
    ]

    with (
        patch("src.features.document_processing.loaders.docxloader.docx_to_pdf", return_value=pdf_path),
        patch.object(PdfLoader, "load", return_value=pdf_blocks),
    ):
        blocks = DocxLoader().load(docx_path)

    assert [block.page for block in blocks] == [1, 2]
    assert [block.line_number for block in blocks] == [1, 1]
    assert all(block.metadata["layout_source"] == "libreoffice_pdf" for block in blocks)


def test_pdf_loader_extracts_blocks_without_table_duplication(tmp_path: Path) -> None:
    pdf_path = tmp_path / "sample.pdf"
    _build_sample_pdf(pdf_path)
    blocks = PdfLoader().load(pdf_path)

    assert blocks
    page_lines = _line_numbers_per_page(blocks)
    assert page_lines[1][0] == 1
    assert page_lines[1] == list(range(1, len(page_lines[1]) + 1))

    table_blocks = [block for block in blocks if block.block_type == "table"]
    text_blocks = [block for block in blocks if block.block_type == "text"]
    assert len(table_blocks) == 1
    assert "Voltage | 230V" in table_blocks[0].text

    flattened = " ".join(block.text for block in text_blocks)
    assert "Voltage230V" not in flattened.replace(" ", "")
    assert flattened.count("Voltage") <= 1

    components = {block_component_type(block) for block in blocks}
    assert COMPONENT_TITLE in components
    assert COMPONENT_SUBTITLE in components
    assert COMPONENT_TABLE in components


def test_pdf_loader_classifies_footer_and_retains_positions(tmp_path: Path) -> None:
    pdf_path = tmp_path / "footer.pdf"
    _build_pdf_with_canvas_footer(pdf_path)
    blocks = PdfLoader().load(pdf_path)
    footer_blocks = [block for block in blocks if block_component_type(block) == COMPONENT_FOOTER]
    assert footer_blocks
    assert footer_blocks[-1].text == "Page 1 of 1"
    assert footer_blocks[-1].line_number == max(block.line_number for block in blocks)


@pytest.mark.skipif(
    not Path(
        r"c:\Users\Adhisivan Ragunathan\OneDrive - Rudhra info solution\Shared\Cleaned\Machinery & Equipments\AUTOCLAVE\HPHV Steam Sterilizer URS (AutoClave).docx"
    ).is_file(),
    reason="URS sample DOCX not available on this machine",
)
def test_urs_docx_line_numbers_are_monotonic_per_page() -> None:
    urs = Path(
        r"c:\Users\Adhisivan Ragunathan\OneDrive - Rudhra info solution\Shared\Cleaned\Machinery & Equipments\AUTOCLAVE\HPHV Steam Sterilizer URS (AutoClave).docx"
    )
    blocks, _ = load_document_blocks(urs)
    page_lines = _line_numbers_per_page(blocks)
    for page, lines in page_lines.items():
        assert lines == sorted(lines)
        assert lines[0] >= 1
        assert len(set(lines)) == len(lines)
