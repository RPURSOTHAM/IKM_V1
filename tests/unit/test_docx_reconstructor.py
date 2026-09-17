import pytest
pytest.skip("legacy docx_template/pdf_template removed; replaced by template_compliance", allow_module_level=True)

from pathlib import Path
import pytest
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.features.document_processing.extractors.docx_template.extractor import extract_template
from src.features.document_processing.extractors.docx_template.reconstructor import reconstruct_docx

def test_docx_reconstruction_round_trip(tmp_path: Path) -> None:
    original_docx = tmp_path / "original.docx"
    doc = Document()
    
    # 1. Create a paragraph with mixed run formatting and paragraph formatting
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    
    r1 = p.add_run("Normal ")
    r2 = p.add_run("BoldItalic ")
    r2.bold = True
    r2.italic = True
    
    r3 = p.add_run("Colored")
    r3.font.color.rgb = RGBColor(255, 0, 0)
    r3.font.name = "Arial"
    r3.font.size = Pt(14)
    
    # 2. Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[1].cells[0].text = "Data 1"
    table.rows[1].cells[1].text = "Data 2"
    
    doc.save(str(original_docx))
    
    # 3. Extract from original docx
    original_template = extract_template(original_docx)
    
    # 4. Reconstruct to a new docx
    reconstructed_docx = tmp_path / "reconstructed.docx"
    reconstruct_docx(original_template, reconstructed_docx)
    
    # 5. Extract from reconstructed docx
    reconstructed_template = extract_template(reconstructed_docx)
    
    # 6. Validate round-trip correctness
    orig_blocks = original_template["content_blocks"]
    recon_blocks = reconstructed_template["content_blocks"]
    
    assert len(orig_blocks) == len(recon_blocks)
    
    # Paragraph checks
    orig_p = orig_blocks[0]
    recon_p = recon_blocks[0]
    
    assert orig_p["alignment"] == recon_p["alignment"]
    assert pytest.approx(orig_p["space_before"]) == recon_p["space_before"]
    assert pytest.approx(orig_p["space_after"]) == recon_p["space_after"]
    
    # Runs checks
    assert len(orig_p["runs"]) == len(recon_p["runs"])
    assert orig_p["runs"][0]["text"] == recon_p["runs"][0]["text"]
    assert orig_p["runs"][1]["bold"] == recon_p["runs"][1]["bold"]
    assert orig_p["runs"][1]["italic"] == recon_p["runs"][1]["italic"]
    assert orig_p["runs"][2]["font_name"] == recon_p["runs"][2]["font_name"]
    assert pytest.approx(orig_p["runs"][2]["font_size"]) == recon_p["runs"][2]["font_size"]
    assert orig_p["runs"][2]["font_color"] == recon_p["runs"][2]["font_color"]
    
    # Table checks
    orig_tables = original_template["tables"]
    recon_tables = reconstructed_template["tables"]
    
    assert len(orig_tables) == len(recon_tables)
    assert orig_tables[0]["rows"] == recon_tables[0]["rows"]
    assert orig_tables[0]["columns"] == recon_tables[0]["columns"]
    assert orig_tables[0]["header_row"] == recon_tables[0]["header_row"]
    assert orig_tables[0]["row_data"] == recon_tables[0]["row_data"]

