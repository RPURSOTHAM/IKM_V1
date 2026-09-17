import pytest
pytest.skip("legacy docx_template/pdf_template removed; replaced by template_compliance", allow_module_level=True)

from pathlib import Path
import pytest
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.features.document_processing.extractors.docx_template.extractor import extract_template

def test_docx_paragraph_formatting_extraction(tmp_path: Path) -> None:
    docx_path = tmp_path / "test_formatting.docx"
    doc = Document()
    
    # 1. Add a heading with center alignment and specific spacing
    p1 = doc.add_paragraph("Heading with Formatting", style="Heading 1")
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(12)
    p1.paragraph_format.space_after = Pt(6)
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.left_indent = Inches(0.5)
    p1.paragraph_format.right_indent = Inches(0.2)
    p1.paragraph_format.first_line_indent = Pt(-10) # Hanging indent example
    
    # 2. Add a normal paragraph with default/inherited formatting
    p2 = doc.add_paragraph("Normal paragraph text with default formatting.")
    
    doc.save(str(docx_path))
    
    # Run extractor
    template = extract_template(docx_path)
    
    # 3. Check recursive heading tree node
    tree = template["document_tree"]
    assert len(tree) >= 1
    h_node = next((node for node in tree if node.get("type") == "heading" and "Heading" in node.get("text")), None)
    assert h_node is not None
    assert h_node["alignment"] == "center"
    assert pytest.approx(h_node["space_before"]) == 12.0
    assert pytest.approx(h_node["space_after"]) == 6.0
    assert pytest.approx(h_node["line_spacing"]) == 1.15
    assert pytest.approx(h_node["left_indent"]) == 36.0 # 0.5 inches = 36 pt
    assert pytest.approx(h_node["right_indent"]) == 14.4 # 0.2 inches = 14.4 pt
    assert pytest.approx(h_node["first_line_indent"]) == -10.0
    
    # 4. Check content blocks
    blocks = template["content_blocks"]
    assert len(blocks) >= 2
    
    # Verify heading block
    h_block = next((b for b in blocks if b["component_type"] == "heading"), None)
    assert h_block is not None
    assert h_block["alignment"] == "center"
    assert pytest.approx(h_block["space_before"]) == 12.0
    assert pytest.approx(h_block["left_indent"]) == 36.0
    
    # Verify normal paragraph block
    p_block = next((b for b in blocks if b["component_type"] == "paragraph"), None)
    assert p_block is not None
    assert p_block["alignment"] is None # Inherited/default
    assert p_block["space_before"] is None
    assert p_block["line_spacing"] is None

