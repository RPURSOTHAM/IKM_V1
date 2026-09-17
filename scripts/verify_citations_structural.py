"""Structural citation verification without LibreOffice."""
from __future__ import annotations

import csv
import re
import sys
from pathlib import Path

from docx import Document
from src.features.documents.infrastructure.content.page_count import docx_page_count_from_app_properties


def extract_docx_structure(path: Path) -> list[dict]:
    doc = Document(str(path))
    items: list[dict] = []
    for i, para in enumerate(doc.paragraphs):
        text = (para.text or "").strip()
        if not text:
            continue
        style = para.style.name if para.style else ""
        items.append({"index": i, "style": style, "text": text})
    for ti, table in enumerate(doc.tables):
        rows_text: list[str] = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            rows_text.append(" | ".join(cells))
        items.append(
            {
                "index": f"table-{ti}",
                "style": "Table",
                "text": "\n".join(rows_text),
                "table_index": ti,
            }
        )
    return items


def section_number(text: str) -> str | None:
    m = re.match(r"^(\d+(?:\.\d+)*)", text.strip())
    return m.group(1) if m else None


def main() -> int:
    docx_path = Path(r"c:\Users\Neevan\Downloads\SOP-CL-TM-001_Study_Startup.docx")
    csv_path = Path(r"c:\Users\Neevan\Downloads\2026-06-25T07-21_export.csv")

    declared_pages = docx_page_count_from_app_properties(docx_path)
    print(f"Word declared page count: {declared_pages}")

    structure = extract_docx_structure(docx_path)
    print(f"\n=== DOCX structure ({len(structure)} items) ===")
    headings: dict[str, str] = {}
    for item in structure:
        text = item["text"]
        sec = section_number(text.split("\n")[0])
        preview = text[:90].replace("\n", " | ")
        if sec or item["style"].startswith("Heading") or item["style"] == "Table":
            print(f"  [{item['style'][:12]:12s}] {preview}")
        if sec and len(sec) <= 4:  # top-level sections like 1.0, 5.1, 5.2
            headings[sec] = text[:100]

    export_rows: list[dict[str, str]] = []
    with open(csv_path, newline="", encoding="utf-8") as f:
        export_rows = list(csv.DictReader(f))

    print(f"\n=== Citation verification ({len(export_rows)} CSV rows) ===\n")

    issues: list[str] = []
    ok = 0

    for row in export_rows:
        csv_page = int(row.get("page") or 0)
        csv_section = (row.get("section_name") or "").strip()
        csv_text = (row.get("text") or "").strip()
        csv_category = (row.get("category") or "").strip()
        csv_sec_num = section_number(csv_section) or section_number(csv_text)

        # Content presence check
        content_found = False
        matching_heading = None
        for item in structure:
            item_text = item["text"]
            norm_csv = re.sub(r"\s+", " ", csv_text.lower())[:60]
            norm_item = re.sub(r"\s+", " ", item_text.lower())
            if norm_csv[:40] in norm_item or norm_item[:40] in norm_csv:
                content_found = True
            if csv_sec_num and item_text.strip().startswith(csv_sec_num):
                matching_heading = item_text[:80]

        # Section label check
        section_ok = True
        if csv_sec_num:
            # For table chunks labeled "5.2 SIV" the section number should be 5.2
            sec_from_label = section_number(csv_section)
            if sec_from_label and sec_from_label != csv_sec_num:
                section_ok = False
            # Check heading exists in doc
            base_sec = csv_sec_num.split(".")[0] + "." + csv_sec_num.split(".")[1] if csv_sec_num.count(".") >= 1 else csv_sec_num
            if csv_sec_num not in headings and base_sec not in headings:
                # might be subsection like 5.1 without dedicated heading line
                found_in_doc = any(csv_sec_num in item["text"] for item in structure)
                if not found_in_doc and not content_found:
                    section_ok = False

        # Page plausibility
        page_ok = True
        if declared_pages and (csv_page < 1 or csv_page > declared_pages):
            page_ok = False

        # Table category check
        category_ok = True
        is_table_content = "|" in csv_text and csv_category == "default" and len(csv_text) > 100
        if is_table_content:
            category_ok = False

        checks = []
        if not content_found:
            checks.append("CONTENT NOT IN DOCX")
        if not section_ok:
            checks.append("SECTION LABEL ISSUE")
        if not page_ok:
            checks.append(f"PAGE OUT OF RANGE (1-{declared_pages})")
        if not category_ok:
            checks.append("TABLE NOT TAGGED AS table")

        if checks:
            issues.append(
                f"FAIL page={csv_page} section={csv_section!r} category={csv_category}\n"
                f"      checks: {', '.join(checks)}\n"
                f"      text: {csv_text[:70]}..."
            )
        else:
            ok += 1
            heading_note = f" (heading: {matching_heading[:50]}...)" if matching_heading else ""
            print(f"OK   page={csv_page} section={csv_section} category={csv_category}{heading_note}")

    if issues:
        print("\n--- Issues ---")
        for issue in issues:
            print(issue)

    # Page distribution analysis
    print(f"\n=== Page distribution in CSV ===")
    pages_used = sorted({int(r["page"]) for r in export_rows})
    print(f"Pages cited: {pages_used} (document has {declared_pages} pages)")

    by_page: dict[int, list[str]] = {}
    for row in export_rows:
        p = int(row["page"])
        by_page.setdefault(p, []).append(row["section_name"])

    for p in pages_used:
        sections = by_page[p]
        print(f"  Page {p}: {sections}")

    # Cross-page consistency for same section
    print(f"\n=== Same-section page consistency ===")
    sec_pages: dict[str, set[int]] = {}
    for row in export_rows:
        sec = (row.get("section_name") or "").strip()
        sec_num = section_number(sec) or sec
        sec_pages.setdefault(sec_num, set()).add(int(row["page"]))

    for sec, pages in sorted(sec_pages.items()):
        if len(pages) > 1:
            print(f"  Section {sec!r} spans pages {sorted(pages)} — {'OK (multi-page section)' if max(pages) - min(pages) <= 1 else 'REVIEW'}")

    print(f"\nSummary: {ok}/{len(export_rows)} rows passed structural checks, {len(issues)} issues")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main())
