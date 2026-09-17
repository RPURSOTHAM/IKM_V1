"""Release Demo — Streamlit UI to test all today-shipped platform features.

Run:
  streamlit run tools/simulators/streamlit_release_demo.py --server.port 8501
"""

from __future__ import annotations

import io
import json
import os
import sys
import time
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import quote

import requests
import streamlit as st

CURRENT_FILE = Path(__file__).resolve()
REPO_ROOT = CURRENT_FILE.parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

try:
    from src.features.chunking.domain.chunking_strategy import CHUNKING_STRATEGIES
except Exception:
    CHUNKING_STRATEGIES = frozenset(
        {
            "fixed-overlap-based",
            "sentence-based",
            "paragraph-based",
            "section-based",
            "hierarchical",
            "semantic",
            "semantic-hierarchy",
            "sliding-window",
        }
    )

DEFAULT_API = os.getenv("RAG_API_BASE_URL", os.getenv("CONSUMER_API_BASE_URL", "http://localhost:8088")).rstrip("/")
DEFAULT_PREFIX = os.getenv("CONSUMER_API_PREFIX", "/api/v1")
DEFAULT_USER = os.getenv("PLATFORM_BOOTSTRAP_ADMIN_USER", "admin")
DEFAULT_PASS = os.getenv("PLATFORM_BOOTSTRAP_ADMIN_PASSWORD", "changeme-local-admin")

APP_TITLE = "RAG Platform — Release Demo"
APP_SUBTITLE = "Swagger-style live testing — pick a feature and call the same DMS APIs."

RELEASE_FEATURES: list[dict[str, Any]] = [
    {
        "id": "ingest",
        "label": "1. Document Management & Ingestion",
        "blurb": "Upload documents into a repository and track intake status.",
        "swagger": ["POST /api/v1/documents/upload", "GET /api/v1/documents", "GET /api/v1/repositories/options"],
    },
    {
        "id": "multiformat",
        "label": "2. Multi-Format Processing & Parsing",
        "blurb": "Process PDF / DOCX / TXT and inspect parsed preview content.",
        "swagger": ["GET /api/v1/documents/{id}/status", "GET /api/v1/documents/{id}/preview", "GET /api/v1/documents/{id}/render"],
    },
    {
        "id": "cih",
        "label": "3. Content Intelligence Hub",
        "blurb": "CIH upload, status, extraction, chunks, and media transcript.",
        "swagger": ["POST /api/v1/cih/upload", "GET /api/v1/cih/{id}/status", "GET /api/v1/cih/{id}/extraction"],
    },
    {
        "id": "chunking",
        "label": "4. Chunking Strategy",
        "blurb": "Browse strategies and run a live chunk preview on custom text.",
        "swagger": [
            "GET /api/v1/chunking/strategies",
            "POST /api/v1/chunking/preview",
            "Upload PDF/DOCX/TXT (local extract → preview)",
        ],
    },
    {
        "id": "embeddings",
        "label": "5. Vector Embeddings Engine",
        "blurb": "Inspect repository embedding settings and document chunk vectors.",
        "swagger": ["GET /api/v1/repositories/{id}", "GET /api/v1/documents/{id}/chunks?include_vector=true"],
    },
    {
        "id": "retrieval",
        "label": "6. Multi-Strategy Retrieval Engine",
        "blurb": "Run retrieval / BM25 search against an indexed repository.",
        "swagger": ["POST /api/v1/retrieve", "POST /api/v1/repositories/{id}/search/bm25"],
    },
    {
        "id": "citations",
        "label": "7. Source Attribution & Citations",
        "blurb": "Ask a grounded question and inspect citation anchors.",
        "swagger": ["POST /api/v1/chat"],
    },
    {
        "id": "references",
        "label": "8. Reference & Template Extraction",
        "blurb": "Look up document references and reverse references.",
        "swagger": ["GET /api/v1/documents/{id}/references", "GET /api/v1/documents/{id}/referenced-by"],
    },
    {
        "id": "rag",
        "label": "9. RAG Generation & LLM Orchestration",
        "blurb": "Full RAG chat with generation models and safety actions.",
        "swagger": ["GET /api/v1/generation/models", "POST /api/v1/chat"],
    },
    {
        "id": "security",
        "label": "10. Security, Safety & DLP",
        "blurb": "Scan text through the security pipeline and prompt-guard health.",
        "swagger": ["POST /api/v1/security/debug/scan", "GET /api/v1/security/prompt-guard"],
    },
    {
        "id": "review",
        "label": "11. Human-in-the-Loop Review",
        "blurb": "List pending security reviews and apply Allow / Mask / Block.",
        "swagger": ["GET /api/v1/security/reviews", "POST /api/v1/security/reviews/{id}/decide"],
    },
    {
        "id": "pubmed",
        "label": "12. PubMed & Clinical Trials",
        "blurb": "Search PubMed + ClinicalTrials.gov evidence, import selected records, and chat with the PubMed AI agent.",
        "swagger": [
            "POST /api/v1/repositories/{id}/evidence/pubmed/search",
            "POST /api/v1/repositories/{id}/evidence/clinical-trials/search",
            "POST /api/v1/repositories/{id}/evidence/import",
            "GET /api/v1/repositories/{id}/evidence/imports",
            "POST /api/v1/repositories/{id}/pubmed-agent/chat",
        ],
    },
    {
        "id": "tenancy",
        "label": "13. Multi-Tenant Repositories",
        "blurb": "List repositories, create one, and inspect tenant options.",
        "swagger": ["GET /api/v1/repositories", "POST /api/v1/repositories", "POST /api/v1/repositories/{id}/activate"],
    },
    {
        "id": "async_jobs",
        "label": "14. Asynchronous Job Processing",
        "blurb": "Monitor document job status and scheduler health.",
        "swagger": ["GET /health", "GET /api/v1/documents/{id}/status", "POST /api/v1/documents/{id}/reprocess"],
    },
    {
        "id": "rendition",
        "label": "15. Rendition Preview for PDF",
        "blurb": "Render PDF pages, PDF view, and download full / selected pages / ZIP.",
        "swagger": [
            "POST /api/v1/rendering/render",
            "GET /api/v1/rendering/{id}/pdf",
            "GET /api/v1/rendering/{id}/pages/{n}",
            "GET /api/v1/rendering/{id}/download?pages=&format=pdf|zip",
        ],
    },
]


# ---------------------------------------------------------------------------
# API helpers
# ---------------------------------------------------------------------------

def _api_base() -> str:
    return str(st.session_state.get("api_base_url") or DEFAULT_API).rstrip("/")


def _api_prefix() -> str:
    return str(st.session_state.get("api_prefix") or DEFAULT_PREFIX).rstrip("/") or "/api/v1"


def _url(path: str) -> str:
    if path.startswith("http"):
        return path
    if path in {"/health", "/status", "/ready", "/metrics"}:
        return f"{_api_base()}{path}"
    if not path.startswith("/"):
        path = f"/{path}"
    if path.startswith(_api_prefix()):
        return f"{_api_base()}{path}"
    return f"{_api_base()}{_api_prefix()}{path}"


def _auth_headers(*, json_body: bool = False) -> dict[str, str]:
    headers: dict[str, str] = {}
    token = str(st.session_state.get("access_token") or "").strip()
    if token:
        headers["Authorization"] = f"Bearer {token}"
    if json_body:
        headers["Content-Type"] = "application/json"
    return headers


def _format_error(exc: Exception) -> str:
    if isinstance(exc, requests.HTTPError) and exc.response is not None:
        try:
            payload = exc.response.json()
        except ValueError:
            return f"HTTP {exc.response.status_code}: {exc.response.text[:500]}"
        detail = payload.get("detail", payload.get("error", payload))
        if isinstance(detail, dict):
            return str(detail.get("message") or detail.get("code") or detail)
        return f"HTTP {exc.response.status_code}: {detail}"
    return str(exc)


def api(
    method: str,
    path: str,
    *,
    timeout: float = 120,
    raw: bool = False,
    **kwargs: Any,
) -> Any:
    headers = dict(kwargs.pop("headers", {}) or {})
    headers.update(_auth_headers(json_body="json" in kwargs))
    full_url = _url(path)
    started = time.perf_counter()
    response = requests.request(
        method,
        full_url,
        timeout=timeout,
        headers=headers or None,
        **kwargs,
    )
    elapsed_ms = int((time.perf_counter() - started) * 1000)
    request_body = kwargs.get("json")
    if request_body is None and kwargs.get("data") is not None:
        request_body = kwargs.get("data")
    last_call = {
        "method": method.upper(),
        "path": path,
        "url": full_url,
        "status_code": response.status_code,
        "elapsed_ms": elapsed_ms,
        "params": kwargs.get("params"),
        "request": request_body,
    }
    try:
        if response.status_code != 204 and (response.text or "").strip() and "application/json" in (
            response.headers.get("content-type") or ""
        ):
            last_call["response"] = response.json()
        else:
            last_call["response_preview"] = (response.text or "")[:800]
    except Exception:
        last_call["response_preview"] = (response.text or "")[:800]
    st.session_state["last_api_call"] = last_call

    if raw:
        return response
    response.raise_for_status()
    if response.status_code == 204 or not (response.text or "").strip():
        return None
    ctype = (response.headers.get("content-type") or "").lower()
    if "application/json" in ctype:
        return response.json()
    return response.content


def ensure_login() -> bool:
    return bool(str(st.session_state.get("access_token") or "").strip())


def do_login(user_id: str, password: str) -> None:
    response = requests.post(
        _url("/auth/token"),
        json={"user_id": user_id, "password": password},
        timeout=30,
    )
    response.raise_for_status()
    payload = response.json()
    token = payload.get("access_token") or payload.get("token")
    if not token:
        raise RuntimeError(f"Token response missing access_token: {payload}")
    st.session_state["access_token"] = token
    st.session_state["auth_user"] = user_id
    st.session_state["auth_payload"] = payload


def auto_login_if_needed() -> None:
    """Sign in automatically for demo (same credentials as Swagger bootstrap admin)."""
    if ensure_login():
        return
    if st.session_state.get("_auto_login_attempted"):
        return
    st.session_state["_auto_login_attempted"] = True
    try:
        user = str(st.session_state.get("login_user") or DEFAULT_USER)
        password = str(st.session_state.get("login_pass") or DEFAULT_PASS)
        do_login(user, password)
    except Exception as exc:
        st.session_state["_auto_login_error"] = str(exc)


def _show_last_api_call() -> None:
    call = st.session_state.get("last_api_call")
    if not isinstance(call, dict):
        return
    with st.expander(
        f"Last API call · {call.get('method')} {call.get('path')} · HTTP {call.get('status_code')} · {call.get('elapsed_ms')} ms",
        expanded=False,
    ):
        st.code(f"{call.get('method')} {call.get('url')}", language="http")
        if call.get("params"):
            st.caption("Query params")
            st.json(call.get("params"))
        if call.get("request") is not None:
            st.caption("Request body")
            st.json(call.get("request"))
        if "response" in call:
            st.caption("Response")
            st.json(call.get("response"))
        elif call.get("response_preview"):
            st.caption("Response preview")
            st.text(call.get("response_preview"))


def _repo_options() -> list[dict[str, Any]]:
    try:
        payload = api("GET", "/repositories/options", timeout=30)
        if isinstance(payload, list):
            return [r for r in payload if isinstance(r, dict)]
    except Exception:
        pass
    try:
        payload = api("GET", "/repositories", timeout=30)
        repos = payload.get("repositories") if isinstance(payload, dict) else payload
        return [r for r in (repos or []) if isinstance(r, dict)]
    except Exception:
        return []


def _repo_label(repo: dict[str, Any]) -> str:
    name = str(repo.get("name") or "").strip()
    rid = str(repo.get("id") or repo.get("repository_id") or "").strip()
    return f"{name} ({rid})" if name else rid


def _repo_id(repo: dict[str, Any]) -> str:
    return str(repo.get("id") or repo.get("repository_id") or "").strip()


def _select_repository(key: str = "demo_repo") -> str | None:
    repos = _repo_options()
    if not repos:
        st.warning("No repositories found. Create one in **13. Multi-Tenant Repositories** first.")
        return None
    labels = {_repo_id(r): _repo_label(r) for r in repos if _repo_id(r)}
    choice = st.selectbox(
        "Repository",
        options=list(labels.keys()),
        format_func=lambda x: labels.get(x, x),
        key=key,
    )
    return choice


def _list_documents(repository_id: str | None = None, limit: int = 50) -> list[dict[str, Any]]:
    params: dict[str, Any] = {"limit": limit, "offset": 0}
    if repository_id:
        params["repository_id"] = repository_id
    payload = api("GET", "/documents", params=params, timeout=60)
    docs = payload.get("documents") if isinstance(payload, dict) else payload
    return [d for d in (docs or []) if isinstance(d, dict)]


def _doc_label(doc: dict[str, Any]) -> str:
    doc_name = str(doc.get("document_name") or "").strip()
    original = str(doc.get("original_file_name") or doc.get("original_filename") or "").strip()
    if doc_name and original and doc_name.lower() != original.lower():
        name = f"{doc_name}  · original: {original}"
    elif original:
        name = f"{original}"
    else:
        name = doc_name or "document"
    status = str(doc.get("status") or "?").strip()
    did = str(doc.get("document_id") or "").strip()
    short = f"{did[:8]}…" if len(did) > 8 else did
    return f"{name} · {status} · {short}"


def _doc_table_rows(docs: list[dict[str, Any]], *, limit: int = 25) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for d in docs[:limit]:
        rows.append(
            {
                "document_id": d.get("document_id"),
                "document_name": d.get("document_name") or "",
                "original_file_name": d.get("original_file_name") or d.get("original_filename") or "",
                "status": d.get("status"),
                "type": d.get("document_type"),
            }
        )
    return rows


def _feature_header(feature: dict[str, Any]) -> None:
    st.markdown(f"## {feature['label']}")
    st.caption(feature["blurb"])
    endpoints = feature.get("swagger") or []
    if endpoints:
        st.markdown("**Swagger endpoints used here**")
        for ep in endpoints:
            st.code(ep, language="http")
    st.divider()


def _require_auth_banner() -> bool:
    if ensure_login():
        return True
    st.error("Sign in from the sidebar first (JWT required for most APIs).")
    return False


# ---------------------------------------------------------------------------
# Feature panels
# ---------------------------------------------------------------------------

def panel_ingest() -> None:
    if not _require_auth_banner():
        return
    repo_id = _select_repository("ingest_repo")
    doc_type_id = None
    if repo_id:
        try:
            type_opts = api("GET", "/document-types/options", params={"repository_id": repo_id}, timeout=30)
            options = {
                str(t.get("id") or t.get("document_type_id")): str(t.get("name") or t.get("id"))
                for t in (type_opts or [])
                if isinstance(t, dict) and (t.get("id") or t.get("document_type_id"))
            }
            if options:
                doc_type_id = st.selectbox(
                    "Document type (optional)",
                    options=["(repository default)", *options.keys()],
                    format_func=lambda x: "Repository default" if x.startswith("(") else options.get(x, x),
                    key="ingest_doc_type",
                )
                if str(doc_type_id).startswith("("):
                    doc_type_id = None
        except Exception as exc:
            st.caption(f"Document types unavailable: {_format_error(exc)}")

    files = st.file_uploader(
        "Upload files (PDF / DOCX / TXT)",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True,
        key="ingest_files",
    )
    submit = st.checkbox("Submit for processing", value=True, key="ingest_submit")
    if st.button("Upload to repository", type="primary", key="ingest_btn"):
        if not repo_id:
            st.error("Select a repository.")
            return
        if not files:
            st.error("Choose at least one file.")
            return
        multipart = []
        for f in files:
            multipart.append(("files", (f.name, f.getvalue(), f.type or "application/octet-stream")))
        data = {
            "repository_id": repo_id,
            "submit_for_processing": str(bool(submit)).lower(),
        }
        if doc_type_id:
            data["document_type_id"] = str(doc_type_id)
        try:
            resp = requests.post(
                _url("/documents/upload"),
                headers=_auth_headers(),
                data=data,
                files=multipart,
                timeout=180,
            )
            st.session_state["last_api_call"] = {
                "method": "POST",
                "path": "/documents/upload",
                "url": _url("/documents/upload"),
                "status_code": resp.status_code,
                "elapsed_ms": None,
                "params": None,
                "request": data,
                "response": resp.json() if (resp.text or "").strip() else None,
            }
            resp.raise_for_status()
            payload = resp.json()
            _show_last_api_call()
            st.success(f"Uploaded {payload.get('uploaded_count', '?')} file(s).")
            rejected_count = int(payload.get("rejected_count") or 0)
            if rejected_count:
                reasons = []
                for item in payload.get("rejected") or []:
                    if isinstance(item, dict):
                        reasons.append(str(item.get("code") or "rejected"))
                reason_hint = ", ".join(sorted(set(reasons))) if reasons else "rejected"
                st.warning(
                    f"{rejected_count} file(s) were not uploaded ({reason_hint}). "
                    "Other files in the batch were accepted."
                )
            st.json(payload)
            st.session_state["last_upload_docs"] = payload.get("documents") or []
        except Exception as exc:
            _show_last_api_call()
            st.error(_format_error(exc))

    st.subheader("Recent documents")
    try:
        docs = _list_documents(repo_id)
        if docs:
            st.dataframe(_doc_table_rows(docs), use_container_width=True)
        else:
            st.info("No documents yet in this repository.")
    except Exception as exc:
        st.error(_format_error(exc))


def panel_multiformat() -> None:
    if not _require_auth_banner():
        return
    st.write("Pick an ingested document and inspect status + read-only preview (parsed content).")
    repo_id = _select_repository("mf_repo")
    docs = []
    try:
        docs = _list_documents(repo_id)
    except Exception as exc:
        st.error(_format_error(exc))
        return
    if not docs:
        st.info("Upload a PDF, DOCX, or TXT via feature 1 first.")
        return
    options = {str(d["document_id"]): _doc_label(d) for d in docs if d.get("document_id")}
    doc_id = st.selectbox("Document", options=list(options.keys()), format_func=lambda x: options[x], key="mf_doc")
    cols = st.columns(3)
    with cols[0]:
        if st.button("Refresh status", key="mf_status"):
            try:
                st.json(api("GET", f"/documents/{doc_id}/status", timeout=60))
                _show_last_api_call()
            except Exception as exc:
                _show_last_api_call()
                st.error(_format_error(exc))
    with cols[1]:
        if st.button("Load preview", key="mf_preview_btn"):
            try:
                preview_payload = api("GET", f"/documents/{doc_id}/preview", timeout=120)
                _show_last_api_call()
                if isinstance(preview_payload, dict):
                    st.session_state["mf_preview_payload"] = preview_payload
                    st.session_state.pop("mf_render_payload", None)
                else:
                    st.error(f"Unexpected preview response type: {type(preview_payload).__name__}")
            except Exception as exc:
                _show_last_api_call()
                st.error(_format_error(exc))
    with cols[2]:
        if st.button("Load render", key="mf_render_btn"):
            try:
                resp = requests.get(
                    _url(f"/documents/{doc_id}/render"),
                    headers=_auth_headers(),
                    timeout=120,
                )
                resp.raise_for_status()
                ctype = (resp.headers.get("content-type") or "").lower()
                st.session_state["mf_render_payload"] = {
                    "content_type": ctype,
                    "bytes": resp.content,
                    "doc_id": doc_id,
                    "filename": f"{doc_id}_render.pdf",
                }
                st.session_state.pop("mf_preview_payload", None)
            except Exception as exc:
                st.error(_format_error(exc))

    # Full-width PDF / preview below the action buttons (not inside a narrow column).
    render = st.session_state.get("mf_render_payload")
    if isinstance(render, dict) and render.get("bytes"):
        ctype = str(render.get("content_type") or "")
        st.caption(f"Content-Type: `{ctype}`")
        if "pdf" in ctype:
            _show_pdf_bytes(
                render["bytes"],
                height=1100,
                filename=str(render.get("filename") or "document.pdf"),
            )
        elif "html" in ctype or "text" in ctype:
            text = render["bytes"].decode("utf-8", errors="replace")
            st.components.v1.html(text, height=700, scrolling=True)
        else:
            st.download_button(
                "Download render bytes",
                data=render["bytes"],
                file_name=f"{render.get('doc_id', 'doc')}_render.bin",
            )

    preview = st.session_state.get("mf_preview_payload")
    if isinstance(preview, dict):
        st.write(f"Format: `{preview.get('format')}` · media: `{preview.get('media_type')}`")
        content = preview.get("content")
        if isinstance(content, str) and content.strip():
            st.text_area("Preview content", content[:12000], height=320)
        else:
            st.json({k: preview.get(k) for k in ("document_id", "document_name", "format", "page_count", "metadata") if k in preview})


def panel_cih() -> None:
    if not _require_auth_banner():
        return
    repo_id = _select_repository("cih_repo")
    files = st.file_uploader(
        "CIH files (PDF / DOCX / PPT / audio / video / image)",
        accept_multiple_files=True,
        key="cih_files",
    )
    if st.button("CIH Upload", type="primary", key="cih_upload") and repo_id and files:
        multipart = [("files", (f.name, f.getvalue(), f.type or "application/octet-stream")) for f in files]
        try:
            resp = requests.post(
                _url("/cih/upload"),
                headers=_auth_headers(),
                data={"repository_id": repo_id, "submit_for_processing": "true"},
                files=multipart,
                timeout=180,
            )
            resp.raise_for_status()
            payload = resp.json()
            st.success("CIH upload accepted.")
            st.json(payload)
            docs = payload.get("documents") or []
            if docs:
                st.session_state["cih_doc_id"] = docs[0].get("document_id")
        except Exception as exc:
            st.error(_format_error(exc))

    doc_id = st.text_input("CIH document_id", value=str(st.session_state.get("cih_doc_id") or ""), key="cih_doc")
    if not doc_id.strip():
        return
    action = st.radio(
        "Inspect",
        ["status", "extraction", "chunks", "metadata", "transcript"],
        horizontal=True,
        key="cih_action",
    )
    if st.button("Fetch CIH result", key="cih_fetch"):
        try:
            st.json(api("GET", f"/cih/{doc_id.strip()}/{action}", timeout=120))
        except Exception as exc:
            st.error(_format_error(exc))


def _extract_text_from_upload(uploaded: Any) -> tuple[str, str]:
    """Extract plain text from an uploaded PDF / DOCX / TXT for chunking preview."""
    filename = str(getattr(uploaded, "name", None) or "upload.txt")
    raw = uploaded.getvalue() if hasattr(uploaded, "getvalue") else bytes(uploaded or b"")
    ext = Path(filename).suffix.lower()

    if ext in {".txt", ".md", ".csv", ".log"}:
        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return raw.decode(encoding), filename
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace"), filename

    if ext == ".pdf":
        try:
            import fitz  # PyMuPDF
        except Exception as exc:
            raise RuntimeError(f"PyMuPDF is required to read PDF uploads: {exc}") from exc
        doc = fitz.open(stream=raw, filetype="pdf")
        try:
            parts: list[str] = []
            for idx, page in enumerate(doc):
                page_text = (page.get_text("text") or "").strip()
                if page_text:
                    parts.append(f"{idx + 1}. Page {idx + 1}\n{page_text}")
            text = "\n\n".join(parts).strip()
            if not text:
                raise RuntimeError("No extractable text found in the PDF.")
            return text, filename
        finally:
            doc.close()

    if ext == ".docx":
        try:
            import docx  # python-docx
        except Exception as exc:
            raise RuntimeError(f"python-docx is required to read DOCX uploads: {exc}") from exc
        document = docx.Document(io.BytesIO(raw))
        parts = [p.text.strip() for p in document.paragraphs if (p.text or "").strip()]
        # Include simple table cell text
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if (c.text or "").strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n\n".join(parts).strip()
        if not text:
            raise RuntimeError("No extractable text found in the DOCX.")
        return text, filename

    raise RuntimeError(f"Unsupported file type '{ext}'. Upload PDF, DOCX, or TXT.")


def panel_chunking() -> None:
    if not _require_auth_banner():
        return
    try:
        catalog = api("GET", "/chunking/strategies", timeout=30)
        strategies = catalog.get("strategies") if isinstance(catalog, dict) else catalog
        if strategies:
            with st.expander("Strategy catalog", expanded=False):
                st.json(strategies)
    except Exception as exc:
        st.warning(f"Could not load catalog: {_format_error(exc)}")

    strategy = st.selectbox(
        "Strategy",
        options=sorted(CHUNKING_STRATEGIES),
        index=sorted(CHUNKING_STRATEGIES).index("section-based")
        if "section-based" in CHUNKING_STRATEGIES
        else 0,
        key="chunk_strategy",
    )

    source = st.radio(
        "Input source",
        ["Upload document", "Paste text", "Existing repository document"],
        horizontal=True,
        key="chunk_source",
    )

    document_name = "preview.txt"
    default_sample = (
        "1. Introduction\nThis SOP describes batch release.\n\n"
        "2. Procedure\nPerform verification steps carefully.\n"
        "Record results in the batch record.\n\n"
        "3. References\nSee SOP-001 and Form-22."
    )

    if source == "Upload document":
        uploaded = st.file_uploader(
            "Upload PDF / DOCX / TXT to chunk",
            type=["pdf", "docx", "txt"],
            key="chunk_upload_file",
        )
        if uploaded is not None:
            fingerprint = f"{uploaded.name}:{len(uploaded.getvalue())}"
            if st.session_state.get("chunk_upload_fingerprint") != fingerprint:
                try:
                    extracted, document_name = _extract_text_from_upload(uploaded)
                    st.session_state["chunk_upload_fingerprint"] = fingerprint
                    st.session_state["chunk_extracted_text"] = extracted
                    st.session_state["chunk_extracted_name"] = document_name
                    st.session_state["chunk_text_from_upload"] = extracted
                    st.success(f"Loaded `{document_name}` · {len(extracted)} characters")
                except Exception as exc:
                    st.error(_format_error(exc))
        text = st.text_area(
            "Extracted text (editable)",
            height=240,
            key="chunk_text_from_upload",
        )
        document_name = str(st.session_state.get("chunk_extracted_name") or document_name)
    elif source == "Existing repository document":
        repo_id = _select_repository("chunk_repo")
        docs = _list_documents(repo_id) if repo_id else []
        if not docs:
            st.info("No documents in this repository. Upload via feature 1 first, or use Upload document.")
            text = ""
        else:
            options = {str(d["document_id"]): _doc_label(d) for d in docs if d.get("document_id")}
            doc_id = st.selectbox(
                "Document",
                list(options.keys()),
                format_func=lambda x: options[x],
                key="chunk_existing_doc",
            )
            selected = next((d for d in docs if str(d.get("document_id")) == doc_id), {})
            document_name = str(
                selected.get("original_file_name")
                or selected.get("document_name")
                or f"{doc_id}.txt"
            )
            if st.button("Load document text", key="chunk_load_doc_text"):
                try:
                    # Prefer preview content; fall back to joining chunks.
                    preview = api("GET", f"/documents/{doc_id}/preview", timeout=120)
                    content = ""
                    if isinstance(preview, dict):
                        content = str(preview.get("content") or "").strip()
                    if not content:
                        chunks_payload = api(
                            "GET",
                            f"/documents/{doc_id}/chunks",
                            params={"limit": 500, "offset": 0, "include_text": True, "include_vector": False},
                            timeout=60,
                        )
                        chunks = (chunks_payload or {}).get("chunks") if isinstance(chunks_payload, dict) else []
                        content = "\n\n".join(
                            str(c.get("text") or "").strip()
                            for c in (chunks or [])
                            if str(c.get("text") or "").strip()
                        )
                    if not content:
                        raise RuntimeError("No text available from preview or chunks for this document.")
                    st.session_state["chunk_extracted_text"] = content
                    st.session_state["chunk_extracted_name"] = document_name
                    st.session_state["chunk_text_from_doc"] = content
                    st.success(f"Loaded text from `{document_name}` · {len(content)} characters")
                except Exception as exc:
                    st.error(_format_error(exc))
            text = st.text_area(
                "Document text (editable)",
                height=240,
                key="chunk_text_from_doc",
            )
            document_name = str(st.session_state.get("chunk_extracted_name") or document_name)
    else:
        text = st.text_area("Sample text", value=default_sample, height=200, key="chunk_text")
        document_name = "sample.txt"

    c1, c2, c3 = st.columns(3)
    size = c1.number_input("chunk_size", min_value=50, max_value=4000, value=500, key="chunk_size")
    overlap = c2.number_input("overlap", min_value=0, max_value=1000, value=50, key="chunk_overlap")
    include_cite = c3.checkbox("Citation metadata", value=True, key="chunk_cite")
    if st.button("Preview chunks", type="primary", key="chunk_preview"):
        if not str(text or "").strip():
            st.error("Provide text via upload, paste, or an existing document first.")
            return
        try:
            result = api(
                "POST",
                "/chunking/preview",
                json={
                    "strategy": strategy,
                    "text": text,
                    "chunk_size": int(size),
                    "chunk_overlap": int(overlap),
                    "citation_retainment": bool(include_cite),
                    "document_name": document_name,
                },
                timeout=60,
            )
            _show_last_api_call()
            st.success(f"`{document_name}` → {result.get('chunk_count', 0)} chunks · strategy `{strategy}`")
            for i, chunk in enumerate(result.get("chunks") or []):
                with st.expander(f"Chunk {i + 1}"):
                    st.write(chunk.get("text") or chunk.get("content") or chunk)
                    meta = {
                        k: chunk.get(k)
                        for k in ("page", "section_name", "section_path", "line_start", "line_end", "source_blocks")
                        if chunk.get(k) is not None
                    }
                    if meta:
                        st.json(meta)
        except Exception as exc:
            _show_last_api_call()
            st.error(_format_error(exc))


def panel_embeddings() -> None:
    if not _require_auth_banner():
        return
    repo_id = _select_repository("emb_repo")
    if repo_id and st.button("Load repository settings", key="emb_settings"):
        try:
            st.json(api("GET", f"/repositories/{repo_id}", timeout=30))
        except Exception as exc:
            st.error(_format_error(exc))
    try:
        models = api("GET", "/generation/models", timeout=30)
        with st.expander("Generation / model catalog"):
            st.json(models)
    except Exception as exc:
        st.caption(f"Model catalog unavailable: {_format_error(exc)}")

    docs = _list_documents(repo_id) if repo_id else []
    if not docs:
        st.info("Index a document first, then inspect chunk vectors here.")
        return
    options = {str(d["document_id"]): _doc_label(d) for d in docs if d.get("document_id")}
    doc_id = st.selectbox("Document", list(options.keys()), format_func=lambda x: options[x], key="emb_doc")
    if st.button("Fetch chunks (with vectors if available)", key="emb_chunks"):
        try:
            payload = api(
                "GET",
                f"/documents/{doc_id}/chunks",
                params={"limit": 5, "offset": 0, "include_text": True, "include_vector": True},
                timeout=60,
            )
            chunks = payload.get("chunks") if isinstance(payload, dict) else payload
            st.write(f"Showing up to 5 chunks. Total: {payload.get('count') or payload.get('total') or len(chunks or [])}")
            for ch in (chunks or [])[:5]:
                vec = ch.get("vector") or ch.get("embedding")
                preview = str(ch.get("text") or "")[:240]
                st.markdown(f"**chunk** `{ch.get('chunk_id') or ch.get('id')}` · dim={len(vec) if isinstance(vec, list) else 'n/a'}")
                st.write(preview)
                if isinstance(vec, list) and vec:
                    st.caption(f"vector[:8] = {vec[:8]}")
        except Exception as exc:
            st.error(_format_error(exc))


def panel_retrieval() -> None:
    if not _require_auth_banner():
        return
    repo_id = _select_repository("ret_repo")
    query = st.text_input("Query", value="What is the purpose of this document?", key="ret_query")
    top_k = st.slider("top_k", 1, 20, 5, key="ret_topk")
    mode = st.radio("Strategy", ["hybrid / retrieve", "bm25"], horizontal=True, key="ret_mode")
    if st.button("Run retrieval", type="primary", key="ret_run") and repo_id:
        try:
            if mode.startswith("hybrid"):
                result = api(
                    "POST",
                    "/retrieve",
                    json={"query": query, "repository_id": repo_id, "top_k": top_k},
                    timeout=90,
                )
            else:
                result = api(
                    "POST",
                    f"/repositories/{repo_id}/search/bm25",
                    json={"query": query, "top_k": top_k},
                    timeout=90,
                )
            st.json(result if isinstance(result, dict) and len(str(result)) < 8000 else {"summary": "ok", "keys": list((result or {}).keys()) if isinstance(result, dict) else type(result).__name__})
            results = []
            if isinstance(result, dict):
                results = result.get("results") or result.get("hits") or result.get("chunks") or []
            for i, item in enumerate(results[:top_k]):
                if isinstance(item, dict):
                    text = item.get("text") or item.get("content") or item.get("snippet") or ""
                    st.markdown(f"**Hit {i + 1}** score=`{item.get('score')}` page=`{item.get('page')}`")
                    st.write(str(text)[:500])
                else:
                    st.write(item)
        except Exception as exc:
            st.error(_format_error(exc))


def _render_citations(citations: list[dict[str, Any]]) -> None:
    if not citations:
        st.info("No citations returned.")
        return
    st.write("### Citations")
    st.caption(
        "`section` comes from section-based chunking. If it is empty, re-ingest with "
        "strategy **section-based** (feature 4 preview / processing config) then ask again."
    )
    for c in citations:
        label_parts = [
            str(c.get("document_name") or c.get("document") or "doc"),
            str(c.get("section") or c.get("section_name") or ""),
            f"page {c.get('page')}" if c.get("page") is not None else "",
        ]
        label = " | ".join(p for p in label_parts if p)
        with st.expander(f"[{c.get('index', '?')}] {label}", expanded=True):
            st.write(c.get("snippet") or c.get("citation_text") or "(no snippet)")
            # Always surface section fields so they are visible in JSON.
            section_view = {
                "section": c.get("section") or c.get("section_name") or "",
                "section_path": c.get("section_path") or "",
                "parent_section": c.get("parent_section") or "",
                "title": c.get("title") or "",
                "page": c.get("page"),
                "page_start": c.get("page_start"),
                "page_end": c.get("page_end"),
                "line_start": c.get("line_start"),
                "line_end": c.get("line_end"),
                "strategy": c.get("strategy") or c.get("strategy_name") or "",
                "document_id": c.get("document_id"),
                "document_name": c.get("document_name") or c.get("document"),
                "citation_text": c.get("citation_text"),
            }
            if c.get("source_blocks") is not None:
                section_view["source_blocks"] = c.get("source_blocks")
            st.json(section_view)
            with st.expander("Full citation object"):
                st.json(c)


def panel_citations() -> None:
    if not _require_auth_banner():
        return
    repo_id = _select_repository("cite_repo")
    query = st.text_input("Question (grounded answer + citations)", key="cite_q")
    if st.button("Ask with citations", type="primary", key="cite_ask") and repo_id and query.strip():
        try:
            chat = api("POST", "/chat", json={"query": query, "repository_id": repo_id}, timeout=90)
            st.session_state["cite_last"] = chat
        except Exception as exc:
            st.error(_format_error(exc))
            return
    chat = st.session_state.get("cite_last")
    if not chat:
        return
    if not chat.get("allowed", True):
        st.error(chat.get("reason") or "Blocked")
        return
    st.write(chat.get("answer"))
    _render_citations(list(chat.get("citations") or []))


def panel_references() -> None:
    if not _require_auth_banner():
        return
    repo_id = _select_repository("ref_repo")
    docs = _list_documents(repo_id) if repo_id else []
    if not docs:
        st.info("No documents available.")
        return
    options = {str(d["document_id"]): _doc_label(d) for d in docs if d.get("document_id")}
    doc_id = st.selectbox("Document", list(options.keys()), format_func=lambda x: options[x], key="ref_doc")
    kind = st.radio("Lookup", ["references", "referenced-by", "reference-graph"], horizontal=True, key="ref_kind")
    ref_base = st.text_input(
        "Reference processor base (optional fallback)",
        value=os.getenv("REFERENCE_EXTRACTION_URL", "http://localhost:3110"),
        key="ref_proc_url",
    )
    if st.button("Load references", type="primary", key="ref_load"):
        path = f"/documents/{quote(doc_id, safe='')}/{kind}"
        try:
            st.json(api("GET", path, timeout=60))
            return
        except Exception as primary_exc:
            st.warning(f"DMS lookup failed: {_format_error(primary_exc)}")
        # Fall back to reference processor if mounted separately
        try:
            url = f"{ref_base.rstrip('/')}/api/v1{path}"
            resp = requests.get(url, headers=_auth_headers(), timeout=60)
            resp.raise_for_status()
            st.json(resp.json())
        except Exception as exc:
            st.error(_format_error(exc))


def panel_rag() -> None:
    if not _require_auth_banner():
        return
    repo_id = _select_repository("rag_repo")
    try:
        models = api("GET", "/generation/models", timeout=30)
        with st.expander("Available generation models"):
            st.json(models)
    except Exception:
        pass
    query = st.text_area("Question", key="rag_q", height=100)
    if st.button("Generate answer", type="primary", key="rag_go") and repo_id and query.strip():
        try:
            chat = api("POST", "/chat", json={"query": query, "repository_id": repo_id}, timeout=120)
            st.session_state["rag_last"] = chat
        except Exception as exc:
            st.error(_format_error(exc))
            return
    chat = st.session_state.get("rag_last")
    if not chat:
        return
    st.write("### Answer")
    st.write(chat.get("answer"))
    st.caption(
        f"action=`{chat.get('action')}` · grounded=`{chat.get('grounded')}` · "
        f"evidence_sufficient=`{chat.get('evidence_sufficient')}` · "
        f"model=`{chat.get('model_id')}` · provider=`{chat.get('provider')}`"
    )
    _render_citations(list(chat.get("citations") or []))
    if chat.get("moderation"):
        with st.expander("Moderation"):
            st.json(chat.get("moderation"))


def panel_security() -> None:
    if not _require_auth_banner():
        return
    cols = st.columns(2)
    with cols[0]:
        if st.button("Prompt-guard diagnostics", key="sec_pg"):
            try:
                st.json(api("GET", "/security/prompt-guard", timeout=30))
            except Exception as exc:
                st.error(_format_error(exc))
    with cols[1]:
        if st.button("Benign self-test", key="sec_benign"):
            try:
                st.json(api("GET", "/security/debug/benign-self-test", timeout=60))
            except Exception as exc:
                st.error(_format_error(exc))

    text = st.text_area(
        "Text to scan",
        value="Hello world test document",
        height=160,
        key="sec_text",
    )
    if st.button("Run security debug scan", type="primary", key="sec_scan"):
        try:
            # Prefer query-style debug endpoint used by DMS
            result = api(
                "POST",
                "/security/debug/scan",
                params={"text": text, "document_id": "streamlit-demo-scan"},
                timeout=90,
            )
            st.json(result)
            status = str(result.get("status") or "").lower()
            if status in {"human_review", "block", "blocked"}:
                st.warning(f"Decision: **{status}** — {result.get('reason')}")
            else:
                st.success(f"Decision: **{status or 'allow'}**")
        except Exception as exc:
            st.error(_format_error(exc))


def panel_review() -> None:
    """End-to-end human-in-the-loop: upload sensitive → human_review → Allow/Mask/Block."""
    if not _require_auth_banner():
        return

    st.markdown(
        """
**How this works**
1. Upload a sensitive file (PDF / DOCX / TXT)  
2. Processing detects DLP / blacklist → status becomes **`human_review`**  
3. Reviewer chooses **ALLOW** · **MASK_AND_ALLOW** · **BLOCK**
"""
    )

    tab_demo, tab_queue = st.tabs(["Demo: upload → review → decide", "Pending review queue"])

    # ------------------------------------------------------------------
    # Tab 1: guided demo
    # ------------------------------------------------------------------
    with tab_demo:
        repo_id = _select_repository("hitl_repo")
        hitl_file = st.file_uploader(
            "Upload sensitive file",
            type=["pdf", "docx", "txt"],
            accept_multiple_files=False,
            key="hitl_file",
        )

        c_up, c_poll = st.columns(2)
        with c_up:
            if st.button("1. Upload sensitive file for processing", type="primary", key="hitl_upload"):
                if not repo_id:
                    st.error("Select a repository first.")
                elif hitl_file is None:
                    st.error("Choose a file to upload.")
                else:
                    try:
                        safe_name = hitl_file.name or "sensitive_hitl_demo.txt"
                        data = {
                            "repository_id": repo_id,
                            "submit_for_processing": "true",
                        }
                        resp = requests.post(
                            _url("/documents/upload"),
                            headers=_auth_headers(),
                            data=data,
                            files=[
                                (
                                    "files",
                                    (
                                        safe_name,
                                        hitl_file.getvalue(),
                                        hitl_file.type or "application/octet-stream",
                                    ),
                                )
                            ],
                            timeout=180,
                        )
                        resp.raise_for_status()
                        payload = resp.json()
                        docs = payload.get("documents") or []
                        doc_id = None
                        review_meta = None
                        if docs and isinstance(docs[0], dict):
                            doc0 = docs[0]
                            doc_id = doc0.get("document_id") or doc0.get("id")
                            meta = doc0.get("metadata") if isinstance(doc0.get("metadata"), dict) else {}
                            scan = meta.get("security_scan") if isinstance(meta.get("security_scan"), dict) else {}
                            review_meta = {
                                "document_id": doc_id,
                                "document_name": doc0.get("document_name") or doc0.get("original_file_name"),
                                "status": doc0.get("status") or scan.get("status"),
                                "review_id": meta.get("review_id") or scan.get("review_id"),
                                "risk_level": scan.get("risk_level") or scan.get("severity"),
                                "severity": scan.get("severity"),
                                "reason": scan.get("reason"),
                                "dlp_decision": scan.get("dlp_decision"),
                                "available_actions": scan.get("available_reviewer_actions")
                                or ["ALLOW", "MASK_AND_ALLOW", "BLOCK"],
                                "detections": scan.get("detections") or [],
                                "detected_categories": scan.get("detected_categories") or [],
                            }
                        if not doc_id and isinstance(payload.get("document_ids"), list) and payload["document_ids"]:
                            doc_id = payload["document_ids"][0]
                        st.session_state["hitl_doc_id"] = str(doc_id or "")
                        st.session_state["hitl_upload_payload"] = payload
                        status_from_upload = str((review_meta or {}).get("status") or "").strip()
                        st.session_state["hitl_status"] = {
                            "status": status_from_upload or None,
                            "document_id": doc_id,
                        }
                        # If upload already escalated, load the review row (or use scan metadata).
                        rid = str((review_meta or {}).get("review_id") or "").strip()
                        if rid:
                            try:
                                review_row = api("GET", f"/security/reviews/{quote(rid)}", timeout=30)
                                if isinstance(review_row, dict):
                                    st.session_state["hitl_matched_review"] = review_row
                                else:
                                    st.session_state["hitl_matched_review"] = review_meta
                            except Exception:
                                st.session_state["hitl_matched_review"] = review_meta
                        else:
                            st.session_state.pop("hitl_matched_review", None)
                        status_now = str((review_meta or {}).get("status") or "").lower()
                        if status_now in {"human_review", "pending_human_review"}:
                            st.warning(
                                f"Uploaded → **human_review** (review `{rid}`). "
                                "Use the decision buttons below."
                            )
                        else:
                            st.success(f"Uploaded. Tracking document `{doc_id}` (status: {status_now or 'unknown'}).")
                        st.json(payload)
                        _show_last_api_call()
                    except Exception as exc:
                        st.error(_format_error(exc))

        doc_id = str(st.session_state.get("hitl_doc_id") or "").strip()
        if doc_id:
            st.info(f"Tracked document: `{doc_id}`")
            with c_poll:
                if st.button("2. Refresh status (wait for human_review)", key="hitl_poll"):
                    try:
                        status_payload = api("GET", f"/documents/{doc_id}/status", timeout=60)
                        st.session_state["hitl_status"] = status_payload
                        # Also try to find matching pending review for this document
                        reviews_payload = api(
                            "GET",
                            "/security/reviews",
                            params={"review_status": "PENDING"},
                            timeout=30,
                        )
                        items = []
                        if isinstance(reviews_payload, dict):
                            items = (
                                reviews_payload.get("reviews")
                                or reviews_payload.get("items")
                                or []
                            )
                        matched = next(
                            (
                                r
                                for r in items
                                if isinstance(r, dict)
                                and str(r.get("document_id") or "") == doc_id
                            ),
                            None,
                        )
                        st.session_state["hitl_matched_review"] = matched
                        _show_last_api_call()
                    except Exception as exc:
                        st.error(_format_error(exc))

            status_payload = st.session_state.get("hitl_status")
            if isinstance(status_payload, dict):
                status = str(
                    status_payload.get("status")
                    or status_payload.get("current_status")
                    or status_payload.get("document_status")
                    or ""
                ).lower()
                if status in {"human_review", "pending_human_review", "in_review"}:
                    st.warning(f"Document status: **{status}** — ready for human review.")
                elif status in {"completed", "complete", "indexed"}:
                    st.success(f"Document status: **{status}** (already past review, or allowed earlier).")
                elif status in {"blocked", "block", "rejected"}:
                    st.error(f"Document status: **{status}**")
                else:
                    st.info(f"Document status: **{status or 'unknown'}** — click refresh again in a few seconds.")
                with st.expander("Status payload"):
                    st.json(status_payload)

            matched = st.session_state.get("hitl_matched_review")
            if isinstance(matched, dict) and (matched.get("review_id") or matched.get("id")):
                st.subheader("3. Human review decision")
                st.caption(
                    f"Review `{matched.get('review_id') or matched.get('id')}` · "
                    f"risk `{matched.get('risk_level') or matched.get('severity')}` · "
                    f"{matched.get('reason') or matched.get('dlp_reason') or ''}"
                )
                detections = matched.get("detections") or []
                if detections:
                    st.write("Detections")
                    st.dataframe(detections, use_container_width=True)
                with st.expander("Full review record"):
                    st.json(matched)

                actions = list(
                    matched.get("available_actions")
                    or ["ALLOW", "MASK_AND_ALLOW", "BLOCK"]
                )
                # Normalize action labels for buttons
                action_labels = {
                    "ALLOW": "ALLOW — continue processing",
                    "MASK_AND_ALLOW": "MASK_AND_ALLOW — mask secrets then continue",
                    "BLOCK": "BLOCK — stop processing",
                }
                reviewer = st.text_input(
                    "Reviewer",
                    value=str(st.session_state.get("auth_user") or DEFAULT_USER),
                    key="hitl_reviewer",
                )
                comments = st.text_input("Comments", value="HITL demo decision", key="hitl_comments")

                b1, b2, b3 = st.columns(3)
                button_cols = {"ALLOW": b1, "MASK_AND_ALLOW": b2, "BLOCK": b3}
                review_id = str(matched.get("review_id") or matched.get("id"))
                for action in ("ALLOW", "MASK_AND_ALLOW", "BLOCK"):
                    if action not in {str(a).upper() for a in actions}:
                        continue
                    col = button_cols[action]
                    with col:
                        if st.button(action_labels.get(action, action), key=f"hitl_decide_{action}", use_container_width=True):
                            try:
                                updated = api(
                                    "POST",
                                    f"/security/reviews/{quote(review_id)}/decide",
                                    json={
                                        "reviewer": reviewer,
                                        "decision": action,
                                        "comments": comments,
                                    },
                                    timeout=60,
                                )
                                st.session_state["hitl_decision_result"] = updated
                                st.session_state.pop("hitl_matched_review", None)
                                st.success(f"Applied **{action}**.")
                                st.json(updated)
                                _show_last_api_call()
                                # Refresh document status after decision
                                try:
                                    st.session_state["hitl_status"] = api(
                                        "GET", f"/documents/{doc_id}/status", timeout=60
                                    )
                                except Exception:
                                    pass
                            except Exception as exc:
                                _show_last_api_call()
                                st.error(_format_error(exc))
            elif doc_id and st.session_state.get("hitl_status"):
                st.info(
                    "No PENDING review row matched this document yet. "
                    "Wait a few seconds and click **Refresh status** again. "
                    "Or open the **Pending review queue** tab."
                )

            if st.session_state.get("hitl_decision_result"):
                with st.expander("Last decision response", expanded=False):
                    st.json(st.session_state["hitl_decision_result"])

    # ------------------------------------------------------------------
    # Tab 2: full pending queue (existing behaviour, clearer actions)
    # ------------------------------------------------------------------
    with tab_queue:
        if st.button("Refresh pending reviews", key="rev_refresh"):
            st.session_state.pop("review_queue", None)
        try:
            payload = api("GET", "/security/reviews", params={"review_status": "PENDING"}, timeout=30)
            _show_last_api_call()
            items = []
            if isinstance(payload, dict):
                items = (
                    payload.get("reviews")
                    or payload.get("items")
                    or payload.get("results")
                    or payload.get("data")
                    or []
                )
            elif isinstance(payload, list):
                items = payload
            st.session_state["review_queue"] = items if isinstance(items, list) else []
        except Exception as exc:
            _show_last_api_call()
            st.error(_format_error(exc))
            return

        items = st.session_state.get("review_queue") or []
        st.caption(f"Pending reviews: **{len(items)}**")
        if not items:
            st.info("No pending human reviews. Use the Demo tab to upload sensitive content.")
            return

        labels = {
            str(i.get("review_id") or i.get("id")): (
                f"{i.get('document_name') or i.get('original_file_name') or i.get('document_id') or 'doc'} · "
                f"{i.get('review_status') or i.get('status')} · "
                f"{i.get('risk_level') or i.get('severity') or ''} · "
                f"{i.get('review_id') or i.get('id')}"
            )
            for i in items
            if (i.get("review_id") or i.get("id"))
        }
        review_id = st.selectbox(
            "Pending review",
            list(labels.keys()),
            format_func=lambda x: labels[x],
            key="rev_pick",
        )
        selected = next(
            (i for i in items if str(i.get("review_id") or i.get("id")) == review_id),
            {},
        )
        st.json(
            {
                "reason": selected.get("reason") or selected.get("dlp_reason"),
                "risk_level": selected.get("risk_level") or selected.get("severity"),
                "dlp_decision": selected.get("dlp_decision"),
                "available_actions": selected.get("available_actions"),
                "detections": selected.get("detections"),
                "document_id": selected.get("document_id"),
                "document_name": selected.get("document_name"),
            }
        )
        with st.expander("Full review JSON"):
            st.json(selected)

        available = [str(a).upper() for a in (selected.get("available_actions") or ["ALLOW", "MASK_AND_ALLOW", "BLOCK"])]
        reviewer = st.text_input(
            "Reviewer",
            value=str(st.session_state.get("auth_user") or DEFAULT_USER),
            key="rev_reviewer",
        )
        comments = st.text_input("Comments", value="Demo review decision", key="rev_reason")
        q1, q2, q3 = st.columns(3)
        for action, col in (("ALLOW", q1), ("MASK_AND_ALLOW", q2), ("BLOCK", q3)):
            if action not in available:
                continue
            with col:
                if st.button(action, key=f"rev_queue_{action}", use_container_width=True):
                    try:
                        updated = api(
                            "POST",
                            f"/security/reviews/{quote(review_id)}/decide",
                            json={
                                "reviewer": reviewer,
                                "decision": action,
                                "comments": comments,
                            },
                            timeout=60,
                        )
                        _show_last_api_call()
                        st.success(f"Applied `{action}`.")
                        st.json(updated)
                        st.session_state.pop("review_queue", None)
                    except Exception as exc:
                        _show_last_api_call()
                        st.error(_format_error(exc))


def _render_evidence_results(payload: Any, *, source_key: str) -> None:
    """Show evidence search hits with optional multi-select import."""
    if not isinstance(payload, dict):
        st.json(payload)
        return
    results = payload.get("results") or []
    st.caption(
        f"source=`{payload.get('source')}` · query=`{payload.get('query')}` · "
        f"count=`{payload.get('count')}` / max=`{payload.get('max_results')}`"
    )
    if not results:
        st.warning("No results. Try a simpler query (e.g. `metformin` or `type 2 diabetes`).")
        with st.expander("Raw response"):
            st.json(payload)
        return

    rows = []
    options: dict[str, dict[str, Any]] = {}
    for item in results:
        if not isinstance(item, dict):
            continue
        rid = str(item.get("record_id") or "").strip()
        title = str(item.get("title") or rid)
        label = f"{rid} — {title[:90]}"
        options[label] = item
        rows.append(
            {
                "id": rid,
                "title": title,
                "date": item.get("published_date"),
                "pdf": item.get("pdf_available"),
                "imported": item.get("already_imported"),
                "url": item.get("external_url"),
            }
        )
    st.dataframe(rows, use_container_width=True)

    selected = st.multiselect(
        "Select records to import into this repository",
        options=list(options.keys()),
        key=f"{source_key}_import_pick",
    )
    submit = st.checkbox("Submit for processing after import", value=True, key=f"{source_key}_import_submit")
    if st.button("Import selected", type="primary", key=f"{source_key}_import_btn") and selected:
        repo_id = st.session_state.get(f"{source_key}_repo_for_import") or st.session_state.get("pm_repo_id")
        # Fall back: callers set pm_active_repo
        repo_id = st.session_state.get("pm_active_repo") or repo_id
        if not repo_id:
            st.error("Select a repository first.")
            return
        items = []
        for label in selected:
            item = options[label]
            src = str(item.get("source") or "").strip()
            rid = str(item.get("record_id") or "").strip()
            if src and rid:
                items.append({"source": src, "record_id": rid})
        if not items:
            st.error("Nothing to import.")
            return
        try:
            imported = api(
                "POST",
                f"/repositories/{repo_id}/evidence/import",
                json={"items": items, "submit_for_processing": bool(submit)},
                timeout=180,
            )
            st.success("Import submitted.")
            st.json(imported)
            _show_last_api_call()
        except Exception as exc:
            _show_last_api_call()
            st.error(_format_error(exc))

    with st.expander("Raw response"):
        st.json(payload)


def panel_pubmed() -> None:
    if not _require_auth_banner():
        return
    repo_id = _select_repository("pm_repo")
    st.session_state["pm_active_repo"] = repo_id
    tab_pm, tab_ct, tab_imports, tab_agent = st.tabs(
        ["PubMed search", "Clinical Trials search", "Imports", "PubMed AI agent"]
    )

    with tab_pm:
        st.caption("POST `/api/v1/repositories/{id}/evidence/pubmed/search`")
        q = st.text_input("PubMed query", value="metformin type 2 diabetes", key="pm_q")
        mesh = st.text_input("MeSH terms (optional, comma-separated)", value="", key="pm_mesh")
        max_results = st.slider("Max results", 1, 20, 5, key="pm_max")
        if st.button("Search PubMed", type="primary", key="pm_search") and repo_id:
            try:
                body: dict[str, Any] = {"query": q, "max_results": max_results}
                mesh_terms = [t.strip() for t in mesh.split(",") if t.strip()]
                if mesh_terms:
                    body["mesh_terms"] = mesh_terms
                result = api(
                    "POST",
                    f"/repositories/{repo_id}/evidence/pubmed/search",
                    json=body,
                    timeout=120,
                )
                st.session_state["pm_last_search"] = result
                _show_last_api_call()
            except Exception as exc:
                _show_last_api_call()
                st.error(_format_error(exc))
        if st.session_state.get("pm_last_search") is not None:
            _render_evidence_results(st.session_state["pm_last_search"], source_key="pm")

    with tab_ct:
        st.caption("POST `/api/v1/repositories/{id}/evidence/clinical-trials/search`")
        cond = st.text_input("Condition", value="type 2 diabetes", key="ct_condition")
        title = st.text_input("Title (optional)", value="", key="ct_title")
        nct = st.text_input("NCT ID (optional, e.g. NCT01234567)", value="", key="ct_nct")
        ct_max = st.slider("Max results", 1, 20, 5, key="ct_max")
        if st.button("Search Clinical Trials", type="primary", key="ct_search") and repo_id:
            try:
                body = {"max_results": ct_max}
                if cond.strip():
                    body["condition"] = cond.strip()
                if title.strip():
                    body["title"] = title.strip()
                if nct.strip():
                    body["nct_id"] = nct.strip().upper()
                if not any(body.get(k) for k in ("condition", "title", "nct_id")):
                    st.error("Provide condition, title, or NCT ID.")
                else:
                    result = api(
                        "POST",
                        f"/repositories/{repo_id}/evidence/clinical-trials/search",
                        json=body,
                        timeout=120,
                    )
                    st.session_state["ct_last_search"] = result
                    _show_last_api_call()
            except Exception as exc:
                _show_last_api_call()
                st.error(_format_error(exc))
        if st.session_state.get("ct_last_search") is not None:
            _render_evidence_results(st.session_state["ct_last_search"], source_key="ct")

    with tab_imports:
        st.caption("GET `/api/v1/repositories/{id}/evidence/imports`")
        if st.button("List imports", type="primary", key="ev_imports_btn") and repo_id:
            try:
                listed = api(
                    "GET",
                    f"/repositories/{repo_id}/evidence/imports",
                    timeout=60,
                )
                st.session_state["ev_imports"] = listed
                _show_last_api_call()
            except Exception as exc:
                _show_last_api_call()
                st.error(_format_error(exc))
        listed = st.session_state.get("ev_imports")
        if listed is not None:
            if isinstance(listed, dict) and isinstance(listed.get("items"), list):
                st.dataframe(listed["items"], use_container_width=True)
            elif isinstance(listed, list):
                st.dataframe(listed, use_container_width=True)
            else:
                st.json(listed)

    with tab_agent:
        st.caption("POST `/api/v1/repositories/{id}/pubmed-agent/chat`")
        aq = st.text_area("Ask the PubMed agent", key="pm_agent_q", height=100)
        if st.button("Chat", type="primary", key="pm_agent_go") and repo_id and aq.strip():
            try:
                result = api(
                    "POST",
                    f"/repositories/{repo_id}/pubmed-agent/chat",
                    json={"query": aq},
                    timeout=120,
                )
                st.write(result.get("answer") if isinstance(result, dict) else result)
                if isinstance(result, dict):
                    cites = result.get("citations") or []
                    if cites:
                        _render_citations(list(cites))
                    with st.expander("Raw response"):
                        st.json(result)
                _show_last_api_call()
            except Exception as exc:
                _show_last_api_call()
                st.error(_format_error(exc))


def panel_tenancy() -> None:
    if not _require_auth_banner():
        return
    if st.button("Refresh repositories", key="ten_refresh"):
        st.session_state.pop("ten_repos", None)
    try:
        repos = _repo_options()
        st.session_state["ten_repos"] = repos
        st.dataframe(
            [
                {
                    "id": _repo_id(r),
                    "name": r.get("name"),
                    "default_document_type": r.get("default_document_type") or r.get("default_document_type_id"),
                    "status": r.get("status"),
                }
                for r in repos
            ],
            use_container_width=True,
        )
    except Exception as exc:
        st.error(_format_error(exc))

    st.subheader("Create repository")
    name = st.text_input("Name", value=f"demo-repo-{int(time.time()) % 100000}", key="ten_name")
    owner = st.text_input("Owner user id", value=st.session_state.get("auth_user") or DEFAULT_USER, key="ten_owner")
    if st.button("Create", type="primary", key="ten_create"):
        try:
            created = api(
                "POST",
                "/repositories",
                json={"name": name, "owner_user_id": owner},
                timeout=60,
            )
            st.success("Repository created.")
            st.json(created)
            st.session_state["ten_created_id"] = (
                (created or {}).get("repository_id")
                or (created or {}).get("id")
                or ((created or {}).get("repository") or {}).get("repository_id")
            )
        except Exception as exc:
            st.error(_format_error(exc))

    created_id = str(st.session_state.get("ten_created_id") or "").strip()
    activate_id = st.text_input("Repository id to activate", value=created_id, key="ten_activate_id")
    if st.button("Activate repository", key="ten_activate") and activate_id:
        try:
            st.json(api("POST", f"/repositories/{activate_id}/activate", timeout=60))
            st.success("Activated.")
        except Exception as exc:
            st.error(_format_error(exc))


def panel_async_jobs() -> None:
    if not _require_auth_banner():
        return
    cols = st.columns(3)
    with cols[0]:
        if st.button("API health", key="job_health"):
            try:
                st.json(api("GET", "/health", timeout=15))
            except Exception as exc:
                st.error(_format_error(exc))
    with cols[1]:
        if st.button("Ready", key="job_ready"):
            try:
                st.json(api("GET", "/ready", timeout=15))
            except Exception as exc:
                st.error(_format_error(exc))
    with cols[2]:
        if st.button("Status", key="job_status"):
            try:
                st.json(api("GET", "/status", timeout=15))
            except Exception as exc:
                st.error(_format_error(exc))

    repo_id = _select_repository("job_repo")
    docs = _list_documents(repo_id) if repo_id else []
    if not docs:
        st.info("No documents to poll.")
        return
    options = {str(d["document_id"]): _doc_label(d) for d in docs if d.get("document_id")}
    doc_id = st.selectbox("Document job", list(options.keys()), format_func=lambda x: options[x], key="job_doc")
    auto = st.checkbox("Auto-refresh every 3s", value=False, key="job_auto")
    placeholder = st.empty()
    try:
        status = api("GET", f"/documents/{doc_id}/status", timeout=30)
        placeholder.json(status)
    except Exception as exc:
        placeholder.error(_format_error(exc))
    if auto:
        time.sleep(3)
        st.rerun()

    if st.button("Reprocess document", key="job_reprocess"):
        try:
            st.json(api("POST", f"/documents/{doc_id}/reprocess", timeout=60))
        except Exception as exc:
            st.error(_format_error(exc))


def _fetch_rendition_page_png(rendition_id: str, page_number: int) -> bytes:
    img = requests.get(_url(f"/rendering/{rendition_id}/pages/{int(page_number)}"), timeout=60)
    img.raise_for_status()
    return img.content


def _fetch_rendition_pdf(rendition_id: str) -> tuple[bytes, str]:
    resp = requests.get(_url(f"/rendering/{rendition_id}/pdf"), timeout=120)
    resp.raise_for_status()
    # Prefer filename from Content-Disposition when present
    cd = resp.headers.get("content-disposition") or ""
    filename = "document.pdf"
    if "filename=" in cd:
        filename = cd.split("filename=", 1)[-1].strip().strip('"')
    return resp.content, filename


def _fetch_rendition_page_pdf(rendition_id: str, page_number: int) -> tuple[bytes, str]:
    resp = requests.get(
        _url(f"/rendering/{rendition_id}/pages/{int(page_number)}/pdf"),
        timeout=60,
    )
    resp.raise_for_status()
    cd = resp.headers.get("content-disposition") or ""
    filename = f"page_{page_number}.pdf"
    if "filename=" in cd:
        filename = cd.split("filename=", 1)[-1].strip().strip('"')
    return resp.content, filename


def _show_pdf_bytes(
    pdf_bytes: bytes,
    *,
    height: int = 1100,
    pdf_url: str | None = None,
    filename: str = "document.pdf",
) -> None:
    """Embed a Chrome-like PDF.js viewer (toolbar + thumbnails) on this page."""
    import base64
    import html as html_lib
    import json

    viewer_height = max(900, int(height))
    safe_name = html_lib.escape(filename or "document.pdf")
    # Always load from bytes inside Streamlit's sandboxed component (avoids CORS).
    use_url = False
    source_js = json.dumps(None)
    # Cap inline payload to keep Streamlit HTML component responsive.
    max_inline = 18 * 1024 * 1024
    if len(pdf_bytes) > max_inline:
        st.warning(
            f"PDF is large ({len(pdf_bytes) / (1024 * 1024):.1f} MB). "
            "Use download, or open the API PDF URL in a tab."
        )
        st.download_button(
            "Download PDF",
            data=pdf_bytes,
            file_name=filename or "document.pdf",
            mime="application/pdf",
        )
        if pdf_url:
            st.markdown(f"[Open PDF URL]({pdf_url})")
        return

    b64 = base64.b64encode(pdf_bytes).decode("ascii")

    st.markdown("---")
    st.markdown("#### PDF viewer")
    st.caption("Proper PDF preview (toolbar · pages · zoom) embedded below on this page.")

    html = f"""
<!DOCTYPE html>
<html>
<head>
  <meta charset="utf-8" />
  <script src="https://cdnjs.cloudflare.com/ajax/libs/pdf.js/3.11.174/pdf.min.js"></script>
  <style>
    :root {{
      --bg: #323639;
      --bar: #323639;
      --bar2: #3c4043;
      --border: #5f6368;
      --text: #e8eaed;
      --muted: #9aa0a6;
      --accent: #8ab4f8;
    }}
    * {{ box-sizing: border-box; }}
    html, body {{
      margin: 0; padding: 0; height: 100%;
      background: var(--bg); color: var(--text);
      font-family: "Segoe UI", system-ui, sans-serif;
      overflow: hidden;
    }}
    #app {{
      display: flex; flex-direction: column;
      width: 100%; height: {viewer_height}px;
      border: 1px solid var(--border); border-radius: 8px; overflow: hidden;
    }}
    #toolbar {{
      display: flex; align-items: center; gap: 10px;
      height: 44px; padding: 0 12px;
      background: var(--bar); border-bottom: 1px solid var(--border);
      flex-shrink: 0;
    }}
    #toolbar .title {{
      font-size: 13px; color: var(--muted); max-width: 220px;
      white-space: nowrap; overflow: hidden; text-overflow: ellipsis;
      margin-right: auto;
    }}
    #toolbar button, #toolbar select {{
      background: var(--bar2); color: var(--text);
      border: 1px solid var(--border); border-radius: 4px;
      height: 28px; min-width: 28px; padding: 0 8px; cursor: pointer;
    }}
    #toolbar button:hover {{ border-color: var(--accent); }}
    #toolbar input[type="number"] {{
      width: 48px; height: 28px; text-align: center;
      background: var(--bar2); color: var(--text);
      border: 1px solid var(--border); border-radius: 4px;
    }}
    #toolbar .sep {{ width: 1px; height: 20px; background: var(--border); margin: 0 4px; }}
    #body {{ display: flex; flex: 1; min-height: 0; }}
    #thumbs {{
      width: 140px; flex-shrink: 0; overflow-y: auto;
      background: #2b2e31; border-right: 1px solid var(--border);
      padding: 8px;
    }}
    .thumb {{
      display: block; width: 100%; margin: 0 0 10px 0; padding: 4px;
      background: #1f2123; border: 2px solid transparent; border-radius: 4px;
      cursor: pointer;
    }}
    .thumb.active {{ border-color: var(--accent); }}
    .thumb canvas {{ width: 100%; height: auto; display: block; }}
    .thumb .num {{ text-align: center; font-size: 11px; color: var(--muted); margin-top: 4px; }}
    #main {{
      flex: 1; overflow: auto; background: #525659;
      display: flex; flex-direction: column; align-items: center;
      padding: 16px 0 40px;
    }}
    .page-wrap {{
      margin: 0 0 16px 0; box-shadow: 0 2px 8px rgba(0,0,0,.45);
      background: #fff; line-height: 0;
    }}
    .page-wrap canvas {{ display: block; }}
    #status {{
      padding: 24px; color: var(--muted); font-size: 14px;
    }}
    #status.err {{ color: #f28b82; }}
  </style>
</head>
<body>
  <div id="app">
    <div id="toolbar">
      <div class="title" title="{safe_name}">{safe_name}</div>
      <button id="prev" title="Previous page">‹</button>
      <input id="pageNum" type="number" min="1" value="1" />
      <span style="color:var(--muted);font-size:13px;">/ <span id="pageCount">–</span></span>
      <button id="next" title="Next page">›</button>
      <div class="sep"></div>
      <button id="zoomOut" title="Zoom out">−</button>
      <select id="zoomSel">
        <option value="0.75">75%</option>
        <option value="1" selected>100%</option>
        <option value="1.25">125%</option>
        <option value="1.5">150%</option>
        <option value="2">200%</option>
        <option value="fit">Fit width</option>
      </select>
      <button id="zoomIn" title="Zoom in">+</button>
    </div>
    <div id="body">
      <div id="thumbs"></div>
      <div id="main"><div id="status">Loading PDF…</div></div>
    </div>
  </div>
  <script>
    pdfjsLib.GlobalWorkerOptions.workerSrc =
      "https://cdnjs.cloudflare.com/ajax/libs/pdf.js/3.11.174/pdf.worker.min.js";

    const SOURCE_URL = {source_js};
    const B64 = "{b64}";
    let pdfDoc = null;
    let scale = 1.0;
    let currentPage = 1;
    const main = document.getElementById("main");
    const thumbs = document.getElementById("thumbs");
    const statusEl = document.getElementById("status");

    function setStatus(msg, isErr) {{
      statusEl.className = isErr ? "err" : "";
      statusEl.textContent = msg;
      if (!main.contains(statusEl)) {{
        main.innerHTML = "";
        main.appendChild(statusEl);
      }}
    }}

    async function getData() {{
      if (SOURCE_URL) {{
        return {{ url: SOURCE_URL, withCredentials: false }};
      }}
      const bin = atob(B64);
      const bytes = new Uint8Array(bin.length);
      for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
      return {{ data: bytes }};
    }}

    async function renderPage(pageNum, targetScale) {{
      const page = await pdfDoc.getPage(pageNum);
      let s = targetScale;
      if (s === "fit") {{
        const base = page.getViewport({{ scale: 1 }});
        s = Math.max(0.5, (main.clientWidth - 48) / base.width);
      }}
      const viewport = page.getViewport({{ scale: s }});
      const canvas = document.createElement("canvas");
      const ctx = canvas.getContext("2d");
      canvas.width = viewport.width;
      canvas.height = viewport.height;
      await page.render({{ canvasContext: ctx, viewport }}).promise;
      return {{ canvas, scaleUsed: s }};
    }}

    async function renderAllPages() {{
      main.innerHTML = "";
      const wraps = [];
      for (let i = 1; i <= pdfDoc.numPages; i++) {{
        const wrap = document.createElement("div");
        wrap.className = "page-wrap";
        wrap.id = "page-" + i;
        main.appendChild(wrap);
        wraps.push(wrap);
      }}
      let used = scale;
      for (let i = 1; i <= pdfDoc.numPages; i++) {{
        const {{ canvas, scaleUsed }} = await renderPage(i, scale);
        used = scaleUsed;
        wraps[i - 1].innerHTML = "";
        wraps[i - 1].appendChild(canvas);
      }}
      if (scale === "fit") {{
        // keep select label; numeric zooms stay as-is
      }}
      document.getElementById("pageCount").textContent = String(pdfDoc.numPages);
      document.getElementById("pageNum").max = pdfDoc.numPages;
      highlightThumb(currentPage);
    }}

    async function renderThumbs() {{
      thumbs.innerHTML = "";
      const maxThumbs = Math.min(pdfDoc.numPages, 40);
      for (let i = 1; i <= maxThumbs; i++) {{
        const page = await pdfDoc.getPage(i);
        const viewport = page.getViewport({{ scale: 0.2 }});
        const canvas = document.createElement("canvas");
        const ctx = canvas.getContext("2d");
        canvas.width = viewport.width;
        canvas.height = viewport.height;
        await page.render({{ canvasContext: ctx, viewport }}).promise;
        const box = document.createElement("div");
        box.className = "thumb" + (i === currentPage ? " active" : "");
        box.dataset.page = String(i);
        box.appendChild(canvas);
        const num = document.createElement("div");
        num.className = "num";
        num.textContent = String(i);
        box.appendChild(num);
        box.onclick = () => goToPage(i);
        thumbs.appendChild(box);
      }}
    }}

    function highlightThumb(n) {{
      currentPage = n;
      document.getElementById("pageNum").value = String(n);
      thumbs.querySelectorAll(".thumb").forEach((el) => {{
        el.classList.toggle("active", Number(el.dataset.page) === n);
      }});
      const el = document.getElementById("page-" + n);
      if (el) el.scrollIntoView({{ behavior: "smooth", block: "start" }});
    }}

    function goToPage(n) {{
      n = Math.max(1, Math.min(pdfDoc.numPages, n));
      highlightThumb(n);
    }}

    document.getElementById("prev").onclick = () => goToPage(currentPage - 1);
    document.getElementById("next").onclick = () => goToPage(currentPage + 1);
    document.getElementById("pageNum").onchange = (e) => goToPage(Number(e.target.value) || 1);
    document.getElementById("zoomSel").onchange = async (e) => {{
      const v = e.target.value;
      scale = v === "fit" ? "fit" : Number(v);
      await renderAllPages();
    }};
    document.getElementById("zoomIn").onclick = async () => {{
      const sel = document.getElementById("zoomSel");
      const opts = ["0.75", "1", "1.25", "1.5", "2"];
      let cur = scale === "fit" ? 1 : Number(scale);
      const next = opts.find((o) => Number(o) > cur + 0.01) || "2";
      sel.value = next;
      scale = Number(next);
      await renderAllPages();
    }};
    document.getElementById("zoomOut").onclick = async () => {{
      const sel = document.getElementById("zoomSel");
      const opts = ["2", "1.5", "1.25", "1", "0.75"];
      let cur = scale === "fit" ? 1 : Number(scale);
      const next = opts.find((o) => Number(o) < cur - 0.01) || "0.75";
      sel.value = next;
      scale = Number(next);
      await renderAllPages();
    }};

    (async function init() {{
      try {{
        const src = await getData();
        pdfDoc = await pdfjsLib.getDocument(src).promise;
        document.getElementById("pageCount").textContent = String(pdfDoc.numPages);
        await renderAllPages();
        renderThumbs(); // thumbs can fill in after pages
      }} catch (err) {{
        setStatus("Failed to load PDF: " + (err && err.message ? err.message : err), true);
      }}
    }})();
  </script>
</body>
</html>
"""
    st.components.v1.html(html, height=viewer_height + 16, scrolling=False)


def _render_rendition_gallery(details: dict[str, Any]) -> None:
    """Show rendered PDF page images / native PDF view inline in Streamlit."""
    ren_id = str(details.get("rendition_id") or "").strip()
    if not ren_id:
        return
    page_count = int(details.get("page_count") or 0)
    doc_name = str(details.get("document_name") or "").strip()
    original = str(details.get("original_file_name") or doc_name).strip()
    st.subheader("PDF preview")
    if doc_name or original:
        if doc_name and original and doc_name != original:
            st.caption(f"**{doc_name}** · original file: `{original}` · {page_count} page(s)")
        else:
            st.caption(f"**{doc_name or original}** · {page_count} page(s)")

    if page_count < 1:
        st.warning("No pages available in this rendition.")
        return

    view_mode = st.radio(
        "Preview mode",
        ["All pages", "Single page", "PDF view"],
        horizontal=True,
        key="ren_view_mode",
    )

    if view_mode == "PDF view":
        try:
            pdf_bytes, pdf_name = _fetch_rendition_pdf(ren_id)
            pdf_url = _url(f"/rendering/{ren_id}/pdf")
            st.caption(f"Viewing PDF: `{pdf_name}`")
            # Real PDF viewer embedded full-width at the bottom of this page.
            _show_pdf_bytes(
                pdf_bytes,
                height=1100,
                pdf_url=pdf_url,
                filename=pdf_name,
            )
            st.download_button(
                "Download full PDF",
                data=pdf_bytes,
                file_name=pdf_name,
                mime="application/pdf",
                key="ren_pdf_view_dl",
                use_container_width=True,
            )
        except Exception as exc:
            st.error(_format_error(exc))
            st.info("Create a **new** rendition after this update so the original PDF is preserved for PDF view.")
        return

    if view_mode == "Single page":
        page_num = st.number_input(
            "Page",
            min_value=1,
            max_value=max(page_count, 1),
            value=1,
            key="ren_page_picker",
        )
        try:
            png = _fetch_rendition_page_png(ren_id, int(page_num))
            st.image(png, caption=f"Page {int(page_num)} of {page_count}", use_container_width=True)
            c1, c2 = st.columns(2)
            with c1:
                st.download_button(
                    "Download this page (PNG)",
                    data=png,
                    file_name=f"{Path(original or 'document').stem}_page_{int(page_num)}.png",
                    mime="image/png",
                    key="ren_page_png_dl",
                    use_container_width=True,
                )
            with c2:
                try:
                    page_pdf, page_pdf_name = _fetch_rendition_page_pdf(ren_id, int(page_num))
                    st.download_button(
                        "Download this page (PDF)",
                        data=page_pdf,
                        file_name=page_pdf_name,
                        mime="application/pdf",
                        key="ren_page_pdf_dl",
                        use_container_width=True,
                    )
                except Exception as exc:
                    st.caption(f"Single-page PDF unavailable: {_format_error(exc)}")
        except Exception as exc:
            st.error(_format_error(exc))
        return

    # All pages — show inline gallery (cap very large PDFs for demo UX)
    max_inline = 20
    show_count = min(page_count, max_inline)
    if page_count > max_inline:
        st.info(f"Showing first {max_inline} of {page_count} pages inline. Use Single page / PDF view for the rest.")
    cols_per_row = 2 if show_count > 1 else 1
    for start in range(0, show_count, cols_per_row):
        cols = st.columns(cols_per_row)
        for offset, col in enumerate(cols):
            page_num = start + offset + 1
            if page_num > show_count:
                break
            with col:
                try:
                    png = _fetch_rendition_page_png(ren_id, page_num)
                    st.image(png, caption=f"Page {page_num} of {page_count}", use_container_width=True)
                except Exception as exc:
                    st.error(f"Page {page_num}: {_format_error(exc)}")


def panel_rendition() -> None:
    st.caption("Rendition routes are public (no JWT required), but repository upload needs auth.")
    mode = st.radio(
        "Source",
        ["Upload PDF (preview only)", "Upload PDF into repository + render", "Existing document_id"],
        key="ren_mode",
    )
    repo_id = None
    doc_id = None
    upload = None
    if mode.startswith("Upload PDF into"):
        if not _require_auth_banner():
            return
        repo_id = _select_repository("ren_repo")
        upload = st.file_uploader("PDF", type=["pdf"], key="ren_file_repo")
    elif mode.startswith("Upload PDF (preview"):
        upload = st.file_uploader("PDF", type=["pdf"], key="ren_file_temp")
    else:
        if not _require_auth_banner():
            return
        repo_id = _select_repository("ren_repo_existing")
        docs = _list_documents(repo_id) if repo_id else []
        pdf_docs = [
            d
            for d in docs
            if str(d.get("document_type") or "").lower() == "pdf"
            or str(d.get("original_file_name") or d.get("document_name") or "").lower().endswith(".pdf")
        ]
        if not pdf_docs:
            pdf_docs = docs
        if pdf_docs:
            options = {str(d["document_id"]): _doc_label(d) for d in pdf_docs if d.get("document_id")}
            doc_id = st.selectbox(
                "Document (document name · original file name)",
                list(options.keys()),
                format_func=lambda x: options[x],
                key="ren_doc",
            )
            selected = next((d for d in pdf_docs if str(d.get("document_id")) == doc_id), {})
            if selected:
                st.write(
                    f"**Document name:** `{selected.get('document_name') or '—'}`  \n"
                    f"**Original file name:** `{selected.get('original_file_name') or selected.get('document_name') or '—'}`"
                )

    if st.button("Create rendition", type="primary", key="ren_go"):
        try:
            if mode.startswith("Existing"):
                if not doc_id:
                    st.error("Select a document.")
                    return
                resp = requests.post(
                    _url("/rendering/render"),
                    data={"document_id": doc_id},
                    timeout=180,
                )
            elif mode.startswith("Upload PDF into"):
                if not upload or not repo_id:
                    st.error("Select repository and PDF.")
                    return
                resp = requests.post(
                    _url("/rendering/render"),
                    data={
                        "repository_id": repo_id,
                        "submit_for_processing": "false",
                    },
                    files={"file": (upload.name, upload.getvalue(), "application/pdf")},
                    timeout=180,
                )
            else:
                if not upload:
                    st.error("Choose a PDF.")
                    return
                resp = requests.post(
                    _url("/rendering/render"),
                    files={"file": (upload.name, upload.getvalue(), "application/pdf")},
                    timeout=180,
                )
            resp.raise_for_status()
            details = resp.json()
            # Keep original filename from the upload when API only returns document_name
            if isinstance(details, dict) and upload is not None:
                details.setdefault("original_file_name", upload.name)
            elif isinstance(details, dict) and doc_id:
                selected = next(
                    (d for d in (_list_documents(repo_id) if repo_id else []) if str(d.get("document_id")) == doc_id),
                    {},
                )
                if selected:
                    details.setdefault("original_file_name", selected.get("original_file_name"))
                    details.setdefault("document_name", selected.get("document_name") or details.get("document_name"))
            st.session_state["ren_details"] = details
            st.success(
                f"Rendition `{details.get('rendition_id')}` · "
                f"{details.get('page_count')} pages · {details.get('status')} · "
                f"file `{details.get('original_file_name') or details.get('document_name')}`"
            )
        except Exception as exc:
            st.error(_format_error(exc))

    details = st.session_state.get("ren_details") or {}
    if not isinstance(details, dict) or not details.get("rendition_id"):
        return

    with st.expander("Rendition metadata", expanded=False):
        st.json(details)

    _render_rendition_gallery(details)

    ren_id = str(details.get("rendition_id") or "").strip()
    page_count = int(details.get("page_count") or 0)
    original = str(details.get("original_file_name") or details.get("document_name") or "document.pdf").strip()
    st.subheader("Download")
    d1, d2 = st.columns(2)
    with d1:
        try:
            pdf_bytes, pdf_name = _fetch_rendition_pdf(ren_id)
            st.download_button(
                "Download full PDF",
                data=pdf_bytes,
                file_name=pdf_name or original,
                mime="application/pdf",
                key="ren_dl_full_pdf",
                use_container_width=True,
            )
        except Exception as exc:
            st.caption(f"PDF download unavailable: {_format_error(exc)}")
    with d2:
        page_for_dl = st.number_input(
            "Single page",
            min_value=1,
            max_value=max(page_count, 1),
            value=1,
            key="ren_dl_page_num",
        )
        try:
            png = _fetch_rendition_page_png(ren_id, int(page_for_dl))
            st.download_button(
                "Download page (PNG)",
                data=png,
                file_name=f"{Path(original).stem}_page_{int(page_for_dl)}.png",
                mime="image/png",
                key="ren_dl_page_png",
                use_container_width=True,
            )
        except Exception as exc:
            st.caption(_format_error(exc))
        try:
            page_pdf, page_pdf_name = _fetch_rendition_page_pdf(ren_id, int(page_for_dl))
            st.download_button(
                "Download page (PDF)",
                data=page_pdf,
                file_name=page_pdf_name,
                mime="application/pdf",
                key="ren_dl_page_pdf",
                use_container_width=True,
            )
        except Exception as exc:
            st.caption(f"Page PDF: {_format_error(exc)}")

    st.markdown("##### Selected pages (range or list)")
    st.caption(
        "Examples: `2-9` (pages 2 through 9) · `6,9,12` (only those pages) · "
        "`1,3,5-8` (mix) · `all`"
    )
    pages_sel = st.text_input(
        "Pages",
        value="2-9",
        key="ren_multi_pages",
        help="Range like 2-9, or specific pages like 6,9,12, or all",
    )
    m1, m2 = st.columns(2)
    with m1:
        if st.button("Prepare selected pages PDF", type="primary", key="ren_multi_pdf"):
            try:
                resp = requests.get(
                    _url(f"/rendering/{ren_id}/download"),
                    params={"pages": pages_sel, "format": "pdf"},
                    timeout=180,
                )
                resp.raise_for_status()
                stem = Path(original).stem or "document"
                safe_pages = "".join(ch if ch.isalnum() or ch in "-_," else "_" for ch in (pages_sel or "pages"))
                st.session_state["ren_multi_pdf_bytes"] = resp.content
                st.session_state["ren_multi_pdf_name"] = f"{stem}_pages_{safe_pages}.pdf"
                st.success(f"Ready: {len(resp.content):,} bytes for pages `{pages_sel}`")
            except Exception as exc:
                st.error(_format_error(exc))
        if st.session_state.get("ren_multi_pdf_bytes"):
            st.download_button(
                "Download selected pages (PDF)",
                data=st.session_state["ren_multi_pdf_bytes"],
                file_name=st.session_state.get("ren_multi_pdf_name") or "selected_pages.pdf",
                mime="application/pdf",
                key="ren_multi_pdf_save",
                use_container_width=True,
            )
    with m2:
        if st.button("Prepare ZIP (PNG pages)", key="ren_zip"):
            try:
                z = requests.get(
                    _url(f"/rendering/{ren_id}/download"),
                    params={"pages": pages_sel, "format": "zip"},
                    timeout=120,
                )
                z.raise_for_status()
                st.session_state["ren_zip_bytes"] = z.content
                st.session_state["ren_zip_name"] = f"{ren_id}_pages.zip"
                with zipfile.ZipFile(io.BytesIO(z.content)) as zf:
                    st.caption("ZIP contains: " + ", ".join(zf.namelist()))
            except Exception as exc:
                st.error(_format_error(exc))
        if st.session_state.get("ren_zip_bytes"):
            st.download_button(
                "Download ZIP",
                data=st.session_state["ren_zip_bytes"],
                file_name=st.session_state.get("ren_zip_name") or "pages.zip",
                mime="application/zip",
                key="ren_zip_save",
                use_container_width=True,
            )


PANELS = {
    "ingest": panel_ingest,
    "multiformat": panel_multiformat,
    "cih": panel_cih,
    "chunking": panel_chunking,
    "embeddings": panel_embeddings,
    "retrieval": panel_retrieval,
    "citations": panel_citations,
    "references": panel_references,
    "rag": panel_rag,
    "security": panel_security,
    "review": panel_review,
    "pubmed": panel_pubmed,
    "tenancy": panel_tenancy,
    "async_jobs": panel_async_jobs,
    "rendition": panel_rendition,
}


# ---------------------------------------------------------------------------
# App shell
# ---------------------------------------------------------------------------

def render_sidebar() -> dict[str, Any]:
    st.sidebar.markdown(f"### {APP_TITLE}")
    st.sidebar.caption(APP_SUBTITLE)
    st.sidebar.text_input("API base URL", value=DEFAULT_API, key="api_base_url")
    st.sidebar.text_input("API prefix", value=DEFAULT_PREFIX, key="api_prefix")

    st.sidebar.markdown("#### Sign in")
    st.sidebar.text_input("User", value=DEFAULT_USER, key="login_user")
    st.sidebar.text_input("Password", value=DEFAULT_PASS, type="password", key="login_pass")
    auto_login_if_needed()
    c1, c2 = st.sidebar.columns(2)
    with c1:
        if st.button("Login", use_container_width=True):
            try:
                do_login(st.session_state["login_user"], st.session_state["login_pass"])
                st.sidebar.success("Signed in")
            except Exception as exc:
                st.sidebar.error(_format_error(exc))
    with c2:
        if st.button("Logout", use_container_width=True):
            st.session_state.pop("access_token", None)
            st.session_state.pop("auth_user", None)
            st.session_state["_auto_login_attempted"] = False
            st.sidebar.info("Signed out")

    if ensure_login():
        st.sidebar.success(f"Authenticated as `{st.session_state.get('auth_user')}`")
    else:
        err = st.session_state.get("_auto_login_error")
        st.sidebar.warning("Not signed in" + (f" — {err}" if err else ""))

    if st.sidebar.button("Check /health", use_container_width=True):
        try:
            st.sidebar.json(api("GET", "/health", timeout=10))
        except Exception as exc:
            st.sidebar.error(_format_error(exc))

    docs_url = f"{_api_base()}/docs"
    st.sidebar.markdown(f"[Open Swagger UI]({docs_url})")

    st.sidebar.markdown("#### Features")
    labels = [f["label"] for f in RELEASE_FEATURES]
    selected_label = st.sidebar.radio(
        "Select feature to test",
        options=labels,
        index=0,
        label_visibility="collapsed",
        key="feature_radio",
    )
    feature = next(f for f in RELEASE_FEATURES if f["label"] == selected_label)
    st.sidebar.caption(feature["blurb"])
    return feature


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, layout="wide", initial_sidebar_state="expanded")
    st.markdown(
        """
        <style>
        section[data-testid="stSidebar"] {background: linear-gradient(180deg, #0f172a 0%, #1e293b 100%);}
        section[data-testid="stSidebar"] * {color: #e2e8f0 !important;}
        div[data-testid="stSidebarNav"] {display:none;}
        .block-container {padding-top: 1.2rem;}
        </style>
        """,
        unsafe_allow_html=True,
    )
    feature = render_sidebar()
    _feature_header(feature)
    panel = PANELS.get(feature["id"])
    if panel:
        panel()
    else:
        st.error(f"No panel wired for {feature['id']}")


if __name__ == "__main__":
    main()
