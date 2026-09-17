"""Generic recursive JSON to DOCX converter."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document

DEFAULT_DOCUMENT_TITLE = "ClinicalTrials.gov Study Document"


def _is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    if isinstance(value, (list, dict)) and len(value) == 0:
        return True
    return False


def _format_primitive(value: Any) -> str:
    if isinstance(value, bool):
        return "True" if value else "False"
    return str(value)


def _heading_level(depth: int) -> int:
    """Match reference layout: Heading 2/3/4 only below the document title."""
    if depth <= 0:
        return 2
    if depth == 1:
        return 3
    return 4


def _append_primitive(doc: Document, key: str, value: Any) -> None:
    doc.add_paragraph(f"{key}: {_format_primitive(value)}")


def _append_list(doc: Document, items: list[Any], depth: int) -> None:
    item_index = 0
    for item in items:
        if _is_empty(item):
            continue
        if isinstance(item, dict):
            item_index += 1
            doc.add_heading(f"Item {item_index}", level=_heading_level(depth))
            _append_mapping(doc, item, depth + 1)
        elif isinstance(item, list):
            _append_list(doc, item, depth + 1)
        else:
            doc.add_paragraph(_format_primitive(item), style="List Bullet")


def _append_mapping(doc: Document, data: dict[str, Any], depth: int) -> None:
    for key, value in data.items():
        if _is_empty(value):
            continue
        label = str(key)
        if isinstance(value, dict):
            doc.add_heading(label, level=_heading_level(depth))
            _append_mapping(doc, value, depth + 1)
        elif isinstance(value, list):
            doc.add_heading(label, level=_heading_level(depth))
            _append_list(doc, value, depth)
        else:
            _append_primitive(doc, label, value)


def json_to_docx_bytes(
    data: Any,
    *,
    title: str = DEFAULT_DOCUMENT_TITLE,
) -> bytes:
    """Convert any JSON-compatible structure into a Word document."""
    doc = Document()
    doc.add_heading(title, level=1)
    if isinstance(data, dict):
        _append_mapping(doc, data, depth=0)
    elif isinstance(data, list):
        _append_list(doc, data, depth=0)
    elif not _is_empty(data):
        doc.add_paragraph(_format_primitive(data))
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def docx_paragraph_signatures(document: Document) -> list[tuple[str, str]]:
    """Return (style, text) tuples for test comparisons."""
    signatures: list[tuple[str, str]] = []
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style else ""
        text = paragraph.text.strip()
        if text:
            signatures.append((style, text))
    return signatures
