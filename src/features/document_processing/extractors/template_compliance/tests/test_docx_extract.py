"""Standalone unit tests for template_compliance (DOCX focus)."""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

_HERE = Path(__file__).resolve().parent
_OPTIMIZER = _HERE.parent
_EXTRACTORS = _OPTIMIZER.parent
for _p in (str(_EXTRACTORS), str(_OPTIMIZER)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from docx import Document
from docx.shared import Inches, Pt

from template_compliance import extract_template, empty_template_document
from template_compliance.exceptions import UnsupportedFormatError


def _build_sample_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    # Header with paragraph fields + a small 2-col table (non-Biocon style chrome).
    header = section.header
    header.paragraphs[0].text = "Document No: TMP-001"
    header.add_paragraph("Title: Sample SOP Template")
    hdr_table = header.add_table(rows=2, cols=2, width=Inches(6))
    hdr_table.cell(0, 0).text = "SOP Title"
    hdr_table.cell(0, 1).text = "Sample SOP Template"
    hdr_table.cell(1, 0).text = "SOP No."
    hdr_table.cell(1, 1).text = "TMP-001"
    for run in header.paragraphs[0].runs:
        run.font.name = "Calibri"
        run.font.size = Pt(10)
    footer = section.footer
    footer.paragraphs[0].text = "Confidential work product"
    for run in footer.paragraphs[0].runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)

    doc.add_paragraph("Title: Microbiological Monitoring of Classified Environment")
    doc.add_paragraph("Document No.: TMP-001")
    doc.add_paragraph("Version No: 1.0")
    doc.add_paragraph("Effective Date: 30/04/2026")
    doc.add_paragraph("Review Date: 29/04/2028")
    # First-page 2-col Key | Value (control + non-control companion).
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.cell(0, 0).text = "Document No.:"
    meta_table.cell(0, 1).text = "TMP-001"
    meta_table.cell(1, 0).text = "Department"
    meta_table.cell(1, 1).text = "Quality Control"
    meta_table.cell(2, 0).text = "Facility"
    meta_table.cell(2, 1).text = "Global"
    # Dynamic signature-style grid (headers from the table — not hardcoded in extractor).
    sig = doc.add_table(rows=4, cols=4)
    for i, label in enumerate(["Signatures", "Signatures", "Signatures", "Signatures"]):
        sig.cell(0, i).text = label
    for i, label in enumerate(["Role", "Name", "Department", "Signature and Date"]):
        sig.cell(1, i).text = label
    sig.cell(2, 0).text = "Prepared By"
    sig.cell(2, 1).text = "Alex Author / QUALITY"
    sig.cell(2, 2).text = "Quality"
    sig.cell(2, 3).text = "06-Feb-2026 12:45:08"
    sig.cell(3, 0).text = "Reviewed By"
    sig.cell(3, 1).text = "Blake Reviewer / QUALITY"
    sig.cell(3, 2).text = "Quality"
    sig.cell(3, 3).text = "09-Feb-2026 10:09:31"

    p1 = doc.add_paragraph("1.0 PURPOSE")
    try:
        p1.style = "Heading 1"
    except Exception:
        pass
    purpose_body = doc.add_paragraph(
        "This procedure describes monitoring of classified environments for product quality."
    )
    try:
        from docx.enum.text import WD_LINE_SPACING

        purpose_body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        purpose_body.paragraph_format.line_spacing = 1.15
    except Exception:
        pass
    p2 = doc.add_paragraph("2.0 SCOPE")
    try:
        p2.style = "Heading 1"
    except Exception:
        pass
    p21 = doc.add_paragraph()
    run21 = p21.add_run("2.1 Applicable Areas")
    run21.bold = True
    proc = doc.add_paragraph("3.0 PROCEDURE")
    try:
        proc.style = "Heading 1"
    except Exception:
        pass
    try:
        from docx.enum.text import WD_LINE_SPACING

        proc_line1 = doc.add_paragraph(
            "Operators shall verify room status before starting classified environment monitoring."
        )
        proc_line1.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        proc_line1.paragraph_format.line_spacing = 1.15
        proc_line2 = doc.add_paragraph(
            "Record results immediately after completing each monitoring activity step."
        )
        proc_line2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        proc_line2.paragraph_format.line_spacing = 1.00
    except Exception:
        doc.add_paragraph(
            "Operators shall verify room status before starting classified environment monitoring."
        )
        doc.add_paragraph(
            "Record results immediately after completing each monitoring activity step."
        )
    doc.add_paragraph("NOTE:")
    doc.add_paragraph("NOTE:")
    doc.add_paragraph("3.1.1 Nested detail that must NOT become a section.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    # Body tables after the first H1 must not be treated as page-1 chrome.
    abbr = doc.add_table(rows=2, cols=4)
    abbr.cell(0, 0).text = "4.2 ABBREVIATIONS:"
    abbr.cell(1, 0).text = "APS"
    abbr.cell(1, 1).text = "Aseptic Process Simulation"
    abbr.cell(1, 2).text = "BMR"
    abbr.cell(1, 3).text = "Batch Manufacturing Record"
    criteria = doc.add_table(rows=2, cols=2)
    criteria.cell(0, 0).text = "Test"
    criteria.cell(0, 1).text = "Acceptance Criteria"
    criteria.cell(1, 0).text = "CCIT"
    criteria.cell(1, 1).text = "Shall comply"
    run = doc.add_paragraph().add_run("body")
    run.font.name = "Arial"
    run.font.size = Pt(11)
    doc.save(str(path))


class SchemaTests(unittest.TestCase):
    def test_schema_keys(self) -> None:
        root = empty_template_document(document_name="x.docx", source_format="docx")
        self.assertIn("document", root)
        doc = root["document"]
        for key in (
            "file_name",
            "file_type",
            "metadata",
            "signature_table",
            "header",
            "footer",
            "logo",
            "fonts",
            "toc",
            "page",
            "sections",
            "repeated_fields",
            "total_images",
            "total_tables",
        ):
            self.assertIn(key, doc)
        self.assertNotIn("document_control", doc)
        self.assertNotIn("first_page_tables", doc)
        self.assertNotIn("tables", doc)
        self.assertNotIn("images", doc)
        self.assertIsInstance(doc["header"], list)
        self.assertIsInstance(doc["footer"], list)
        fonts = doc["fonts"]
        self.assertIn("header", fonts)
        self.assertIn("section_heading", fonts)
        self.assertIn("section_content", fonts)


class DocxExtractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._tmpdir = tempfile.TemporaryDirectory()
        cls.docx_path = Path(cls._tmpdir.name) / "sample_template.docx"
        _build_sample_docx(cls.docx_path)
        cls.root = extract_template(cls.docx_path)
        cls.doc = cls.root["document"]

    @classmethod
    def tearDownClass(cls) -> None:
        cls._tmpdir.cleanup()

    def test_shape(self) -> None:
        self.assertEqual(self.doc["file_type"], "docx")
        self.assertIsInstance(self.doc["header"], list)
        self.assertIsInstance(self.doc["footer"], list)
        self.assertIsInstance(self.doc["logo"], dict)
        self.assertEqual(set(self.doc["logo"].keys()), {"present", "location"})
        self.assertIn("present", self.doc["logo"])
        self.assertIsInstance(self.doc["logo"]["present"], bool)
        self.assertIn(str(self.doc["logo"].get("location") or ""), {"", "header", "footer"})
        if self.doc["logo"]["present"]:
            self.assertIn(self.doc["logo"]["location"], {"header", "footer"})
        else:
            self.assertEqual(self.doc["logo"]["location"], "")
        self.assertIsInstance(self.doc["toc"], dict)
        self.assertIn("present", self.doc["toc"])
        self.assertIsInstance(self.doc["page"], dict)
        self.assertIn("present", self.doc["page"])
        self.assertIn("total_images", self.doc)
        self.assertIn("total_tables", self.doc)

    def test_metadata_slots(self) -> None:
        meta = self.doc["metadata"]
        for key in (
            "document_type",
            "title",
            "document_no",
            "version_no",
            "effective_date",
            "review_date",
        ):
            self.assertIn(key, meta)
        self.assertTrue(meta["title"])
        self.assertTrue(meta["document_no"])
        # Non-control first-page table keys must NOT pollute metadata.
        self.assertNotIn("Department", meta)
        self.assertNotIn("Facility", meta)
        self.assertNotIn("Definitions", meta)
        self.assertNotIn("document_control", self.doc)

    def test_signature_table_dynamic_columns(self) -> None:
        sig = self.doc["signature_table"]
        self.assertEqual(sig.get("title"), "Signatures")
        self.assertEqual(
            sig.get("columns"),
            ["Role", "Name", "Department", "Signature and Date"],
        )
        rows = sig.get("rows") or []
        self.assertGreaterEqual(len(rows), 2)
        # Keys are derived from header cells dynamically.
        self.assertIn("role", rows[0])
        self.assertIn("name", rows[0])
        self.assertIn("department", rows[0])
        self.assertIn("signature_and_date", rows[0])
        self.assertEqual(rows[0]["role"], "Prepared By")
        # Signature lives only in signature_table (first_page_tables removed).
        self.assertNotIn("first_page_tables", self.doc)

    def test_signature_skipped_without_heading(self) -> None:
        from template_compliance.docx_tables import _pick_signature_table

        leaked = _pick_signature_table(
            [
                {
                    "kind": "grid",
                    "title": "ABBREVIATION(S)",
                    "columns": ["Abbreviation", "Complete form"],
                    "rows": [{"abbreviation": "SOP", "complete_form": "Standard operating procedure"}],
                },
                {
                    "kind": "grid",
                    "title": "",
                    "columns": ["Class", "Risk Level", "Description"],
                    "rows": [{"class": "Critical", "risk_level": "High", "description": "x" * 40}],
                },
            ]
        )
        self.assertEqual(leaked.get("title"), "")
        self.assertEqual(leaked.get("columns"), [])
        self.assertEqual(leaked.get("rows"), [])

        kept = _pick_signature_table(
            [
                {
                    "kind": "grid",
                    "title": "Signatures",
                    "columns": ["Role", "Name", "Department", "Signature and Date"],
                    "rows": [
                        {
                            "role": "Prepared By",
                            "name": "A",
                            "department": "QA",
                            "signature_and_date": "01-Jan-2026",
                        }
                    ],
                }
            ]
        )
        self.assertEqual(kept.get("title"), "Signatures")
        self.assertEqual(len(kept.get("rows") or []), 1)

    def test_fonts(self) -> None:
        fonts = self.doc["fonts"]
        self.assertIn("name", fonts["header"])
        self.assertIn("size", fonts["header"])
        self.assertIn("name", fonts["footer"])
        self.assertIn("size", fonts["footer"])
        self.assertIn("name", fonts["section_heading"])
        self.assertIn("name", fonts["section_content"])
        spacing = str(fonts.get("minimum_line_spacing") or "")
        self.assertTrue(spacing)
        self.assertNotIn("4.0 pt", spacing)
        self.assertRegex(spacing, r"^\d+\.\d+x$|^\d+(\.\d+)?\s*pt$")

    def test_header_footer_content(self) -> None:
        self.assertTrue(self.doc.get("header"))
        self.assertTrue(any("Document No" in str(x) for x in self.doc["header"]))
        self.assertTrue(self.doc.get("footer"))
        self.assertTrue(any("Confidential" in str(x) for x in self.doc["footer"]))

    def test_chrome_garbage_filter(self) -> None:
        from template_compliance.docx_extract import _clean_chrome_lines

        cleaned = _clean_chrome_lines(
            [
                "ect1ve",
                "ect1veect1ve",
                "Document No: TMP-001",
                "STANDARD OPERATING PROCEDURE",
                "xqzplm",
            ]
        )
        self.assertIn("Document No: TMP-001", cleaned)
        self.assertIn("STANDARD OPERATING PROCEDURE", cleaned)
        self.assertFalse(any("ect1ve" in x.casefold() for x in cleaned))
        self.assertFalse(any(x.casefold() == "xqzplm" for x in cleaned))

    def test_sections(self) -> None:
        sections = self.doc["sections"]
        self.assertTrue(sections)
        for section in sections:
            self.assertNotIn("font", section)
            self.assertIn("image_count", section)
            self.assertIn("table_count", section)
            self.assertIn("indentation", section)
            indent = section["indentation"]
            self.assertIsInstance(indent, dict)
            self.assertIn("indent", indent)
            self.assertIn("alignment", indent)
            self.assertIn(indent.get("alignment") or "", {"", "Left", "Right", "Center", "Justified"})
            self.assertRegex(str(section["number"]), r"^\d+$")
            for child in section.get("subsections") or []:
                self.assertIn("image_count", child)
                self.assertIn("table_count", child)
                self.assertIsInstance(child.get("indentation"), dict)
                self.assertIn("alignment", child["indentation"])
                self.assertRegex(str(child["number"]), r"^\d+\.\d+$")
                self.assertLess(str(child["number"]).count("."), 2)

    def test_repeated_fields(self) -> None:
        self.assertIsInstance(self.doc["repeated_fields"], list)

    def test_outline_title_accepts_title_case_majors(self) -> None:
        from template_compliance.docx_common import (
            _is_banner_chrome_title,
            _is_valid_outline_title,
        )

        self.assertTrue(_is_valid_outline_title("PURPOSE", 1))
        self.assertTrue(_is_valid_outline_title("Objective", 1))
        self.assertTrue(_is_valid_outline_title("CAPA Procedure", 1))
        self.assertTrue(_is_valid_outline_title("Definitions", 1))
        self.assertTrue(_is_valid_outline_title("LIST OF ANNEXURES/ ATTACHMENTS", 1))
        self.assertFalse(
            _is_valid_outline_title("CONFIDENTIAL – FOR INTERNAL USE ONLY", 1)
        )
        self.assertTrue(_is_banner_chrome_title("CONFIDENTIAL – FOR INTERNAL USE ONLY"))
        self.assertFalse(_is_banner_chrome_title("PURPOSE"))
        self.assertFalse(_is_banner_chrome_title("REVISION HISTORY"))

    def test_heading_noise_rejects_ocr_glyphs(self) -> None:
        from template_compliance.docx_common import _is_heading_noise

        self.assertTrue(_is_heading_noise('"C Ql C'))
        self.assertTrue(_is_heading_noise("C Ql C"))
        self.assertFalse(_is_heading_noise("PURPOSE"))
        self.assertFalse(_is_heading_noise("Label for Media Fill Contaminated Units Format"))

    def test_json(self) -> None:
        self.assertIn("document", json.dumps(self.root))


class PdfStubTests(unittest.TestCase):
    def test_pdf_not_implemented(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "x.pdf"
            pdf.write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
            with self.assertRaises(UnsupportedFormatError):
                extract_template(pdf)


if __name__ == "__main__":
    unittest.main()
