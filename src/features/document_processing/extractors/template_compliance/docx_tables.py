"""DOCX extract helper module (split from docx_extract)."""

from __future__ import annotations

import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .page_meta import (
    colon_fields_from_page_lines,
    column_key,
    is_valid_meta_key,
    metadata_from_colon_fields,
    metadata_slot_for_key,
)
from .docx_common import (
    W_NS,
    _colon_fields_from_lines,
    _is_valid_outline_title,
    _outline_level,
    _parse_outline,
    _unique_fields,
)
from .docx_images import (
    _images_from_paragraphs,
)

def _extract_page1_tables_from_docx(
    doc: Document,
    *,
    path: Path | None,
) -> list[dict[str, Any]]:
    """
    Build chrome tables from python-docx Table objects that appear *before*
    the first body H1 (e.g. PURPOSE). Optional page-break indexes can only
    narrow that set — never expand it to mid-document procedure tables.
    """
    total = len(getattr(doc, "tables", []) or [])
    if total <= 0:
        return []

    before_section = _body_table_indexes_before_first_section(doc)
    page1 = _page1_table_indexes(path) if path is not None else None
    if page1 is None:
        indexes = before_section
    else:
        indexes = before_section & page1

    tables: list[dict[str, Any]] = []
    for table_index in sorted(i for i in indexes if 0 <= i < total):
        try:
            grid = _table_grid_from_docx_table(doc.tables[table_index])
        except Exception:
            continue
        if not grid:
            continue
        parsed = _parse_dynamic_table(grid, preceding_title="")
        if parsed:
            parsed["_table_index"] = table_index
            tables.append(parsed)
    return tables


def _table_grid_from_docx_table(table) -> list[list[str]]:
    """Cell grid with merged-cell duplicates collapsed (by underlying tc identity)."""
    grid: list[list[str]] = []
    for row in table.rows:
        cells_out: list[str] = []
        seen_tc: set[int] = set()
        for cell in row.cells:
            try:
                tc_id = id(cell._tc)
            except Exception:
                tc_id = id(cell)
            if tc_id in seen_tc:
                continue
            seen_tc.add(tc_id)
            cells_out.append(" ".join((cell.text or "").split()).strip())
        while cells_out and not cells_out[-1]:
            cells_out.pop()
        if cells_out:
            grid.append(cells_out)
    return grid


def _page1_table_indexes(path: Path | None) -> set[int] | None:
    """
    Return body table indexes that start on page 1 using lastRenderedPageBreak /
    explicit page breaks.

    Returns None when markers are missing/unreliable so callers do not fall back
    to "first N tables" (which pulls Abbreviations / criteria grids on many SOPs).
    """
    if path is None:
        return None
    try:
        with zipfile.ZipFile(path) as archive:
            if "word/document.xml" not in archive.namelist():
                return None
            xml = archive.read("word/document.xml")
    except Exception:
        return None

    try:
        root = ET.fromstring(xml)
    except Exception:
        return None

    body = root.find(f"{{{W_NS}}}body")
    if body is None:
        return None

    page1: set[int] = set()
    table_index = 0
    page = 1
    saw_break = False

    def _node_has_page_break(node: ET.Element) -> bool:
        if node.find(f".//{{{W_NS}}}lastRenderedPageBreak") is not None:
            return True
        for br in node.findall(f".//{{{W_NS}}}br"):
            if str(br.get(f"{{{W_NS}}}type") or "").casefold() == "page":
                return True
        return False

    for child in list(body):
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            if page == 1:
                page1.add(table_index)
            table_index += 1
            # Page break markers sometimes sit inside the table that ends page 1.
            if _node_has_page_break(child):
                saw_break = True
                page = 2
            continue
        if _node_has_page_break(child):
            saw_break = True
            page = 2

    if not saw_break:
        return None
    return page1


def _body_table_indexes_before_first_section(doc: Document) -> set[int]:
    """Top-level ``w:tbl`` indexes that appear before the first body H1."""
    before: set[int] = set()
    table_index = 0
    body = getattr(getattr(doc, "element", None), "body", None)
    if body is None:
        return before

    for child in list(body):
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "tbl":
            before.add(table_index)
            table_index += 1
            continue
        if tag != "p":
            continue
        texts = [
            (node.text or "").strip()
            for node in child.findall(f".//{{{W_NS}}}t")
            if (node.text or "").strip()
        ]
        line = " ".join(texts).strip()
        if line and _is_first_body_section_heading(line):
            break
    return before


def _is_first_body_section_heading(text: str) -> bool:
    """True for the first major body outline heading (e.g. ``1.0 PURPOSE``)."""
    line = " ".join(str(text or "").split()).strip()
    if not line:
        return False
    number, title = _parse_outline(line)
    if number and title and _outline_level(number) == 1 and _is_valid_outline_title(title, 1):
        return True
    return False


def _is_body_section_table(table: dict[str, Any]) -> bool:
    """
    True when a parsed table clearly belongs to numbered body content
    (e.g. title ``4.2 ABBREVIATIONS:``) rather than cover-page chrome.
    """
    title = " ".join(str(table.get("title") or "").split()).strip()
    if re.match(r"^\d+(?:\.\d+)+\b", title):
        return True
    # Key/value grids whose first key looks like a procedure heading.
    if str(table.get("kind") or "") == "key_value":
        pairs = table.get("pairs") or []
        if pairs:
            first_key = " ".join(str((pairs[0] or {}).get("key") or "").split()).strip()
            if re.match(r"^\d+(?:\.\d+)+\b", first_key):
                return True
    # Grid headers that are the first abbrev row (no real header), flagged via title empty
    # and first column looking like symbols / very short tokens across a wide glossary.
    columns = [str(c or "").strip() for c in (table.get("columns") or [])]
    if (
        str(table.get("kind") or "") == "grid"
        and not title
        and len(columns) >= 4
        and sum(1 for c in columns if len(c) <= 3) >= 2
    ):
        return True
    return False


def _materialize_first_page_tables(
    row_cells: dict[tuple[int, int], dict[int, str]],
    table_titles: dict[int, str],
) -> list[dict[str, Any]]:
    """Turn collected cell maps into dynamic grid / key_value table objects."""
    if not row_cells:
        return []
    by_table: dict[int, dict[int, dict[int, str]]] = {}
    for (table_index, row_index), cells in row_cells.items():
        by_table.setdefault(table_index, {})[row_index] = cells

    tables: list[dict[str, Any]] = []
    for table_index in sorted(by_table):
        rows_map = by_table[table_index]
        grid: list[list[str]] = []
        for row_index in sorted(rows_map):
            cells = rows_map[row_index]
            if not cells:
                continue
            max_col = max(cells)
            row = [str(cells.get(i) or "").strip() for i in range(max_col + 1)]
            # Drop trailing empties so chrome tables don't inflate into grids.
            while row and not row[-1]:
                row.pop()
            if row:
                grid.append(row)
        if not grid:
            continue
        parsed = _parse_dynamic_table(
            grid,
            preceding_title=str(table_titles.get(table_index) or ""),
        )
        if parsed:
            parsed["_table_index"] = table_index
            tables.append(parsed)
    return tables


def _dedupe_first_page_tables(tables: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Drop exact structural duplicates (repeating header chrome tables)."""
    out: list[dict[str, Any]] = []
    seen: set[str] = set()
    for table in tables:
        key = repr(
            (
                str(table.get("kind") or ""),
                str(table.get("title") or ""),
                tuple(table.get("columns") or []),
                tuple(
                    tuple(sorted((str(k), str(v)) for k, v in (row or {}).items()))
                    for row in (table.get("rows") or [])
                ),
                tuple(
                    (str(p.get("key") or ""), str(p.get("value") or ""))
                    for p in (table.get("pairs") or [])
                ),
            )
        )
        if key in seen:
            continue
        seen.add(key)
        cleaned = dict(table)
        cleaned.pop("_table_index", None)
        out.append(cleaned)
    return out


def _control_pairs_from_table(table: dict[str, Any]) -> list[dict[str, str]]:
    """Extract control Key/Value pairs from a parsed first-page table."""
    pairs: list[dict[str, str]] = []
    if str(table.get("kind") or "") == "key_value":
        for pair in table.get("pairs") or []:
            name = str(pair.get("key") or "").strip()
            value = str(pair.get("value") or "").strip()
            if name and value and metadata_slot_for_key(name):
                pairs.append({"key": name, "value": value})
        return pairs

    for row in table.get("rows") or []:
        if not isinstance(row, dict):
            continue
        texts = [str(v or "").strip() for v in row.values() if str(v or "").strip()]
        joined = " ".join(texts)
        for item in colon_fields_from_page_lines(texts + ([joined] if joined else [])):
            name = str(item.get("name") or "").strip()
            value = str(item.get("value") or "").strip()
            if name and value and metadata_slot_for_key(name):
                pairs.append({"key": name, "value": value})
    return pairs


def _parse_dynamic_table(
    grid: list[list[str]],
    *,
    preceding_title: str = "",
) -> dict[str, Any] | None:
    """
    Classify a first-page table without hardcoded column/title names.

    - label/value morphology (2 effective columns OR label: cells) → key_value
    - wider grids with a real header row → grid
    """
    if not grid:
        return None
    width = max((len(r) for r in grid), default=0)
    if width <= 0:
        return None

    rows = [list(r) + [""] * (width - len(r)) for r in grid]
    # Trim empty trailing columns for morphology decisions.
    usable_width = 0
    for row in rows:
        for i, cell in enumerate(row):
            if cell.strip():
                usable_width = max(usable_width, i + 1)
    if usable_width <= 0:
        return None
    rows = [row[:usable_width] for row in rows]
    width = usable_width

    # Prefer key/value when left cells look like labels.
    if width == 2 or _looks_like_label_value_table(rows):
        pairs = _pairs_from_label_value_rows(rows)
        if pairs:
            return {
                "kind": "key_value",
                "title": preceding_title.strip(),
                "pairs": pairs,
                "columns": [],
                "rows": [],
            }
        if width == 2:
            return None

    # Grid (≥3 cols): optional title row then header row then data.
    title = ""
    header_idx = 0
    first_nonempty = [c for c in rows[0] if c.strip()]
    if len(rows) >= 2 and first_nonempty:
        unique = {c.strip() for c in first_nonempty}
        if len(unique) == 1 and (len(first_nonempty) >= 2 or width >= 3):
            title = next(iter(unique))
            header_idx = 1
    if not title:
        title = preceding_title.strip()

    if header_idx >= len(rows):
        return None
    columns = [c.strip() for c in rows[header_idx]]
    while columns and not columns[-1]:
        columns.pop()
    if len([c for c in columns if c]) < 2:
        return None
    # Header cells that are themselves "Title:" / "Document No.:" are chrome labels,
    # not a signature grid — convert via label/value pairing instead.
    labelish = sum(1 for c in columns if c.endswith(":") or metadata_slot_for_key(c.rstrip(":")))
    if labelish >= max(1, len([c for c in columns if c]) // 2):
        pairs = _pairs_from_label_value_rows(rows[header_idx:])
        if pairs:
            return {
                "kind": "key_value",
                "title": title,
                "pairs": pairs,
                "columns": [],
                "rows": [],
            }

    keys: list[str] = []
    used: set[str] = set()
    for i, label in enumerate(columns):
        key = column_key(label) if label else ""
        if not key:
            key = f"col_{i + 1}"
        base = key
        n = 2
        while key in used:
            key = f"{base}_{n}"
            n += 1
        used.add(key)
        keys.append(key)

    data_rows: list[dict[str, str]] = []
    for row in rows[header_idx + 1 :]:
        values = [c.strip() for c in row[: len(columns)]]
        if not any(values):
            continue
        if [v.casefold() for v in values] == [c.casefold() for c in columns]:
            continue
        item = {keys[i]: (values[i] if i < len(values) else "") for i in range(len(keys))}
        if any(item.values()):
            data_rows.append(item)

    return {
        "kind": "grid",
        "title": title,
        "columns": columns,
        "rows": data_rows,
        "pairs": [],
    }


def _looks_like_label_value_table(rows: list[list[str]]) -> bool:
    """True when most rows look like Label: | Value morphology."""
    if not rows:
        return False
    hits = 0
    checked = 0
    for row in rows[:8]:
        cells = [c.strip() for c in row if c.strip()]
        if len(cells) < 2:
            continue
        checked += 1
        left = cells[0]
        # Require explicit Label: morphology — do not treat grid headers like
        # "Role" / "Name" as key/value pairs.
        if left.endswith(":") or (":" in left and is_valid_meta_key(left.split(":", 1)[0])):
            hits += 1
    return checked > 0 and hits / checked >= 0.6


def _pairs_from_label_value_rows(rows: list[list[str]]) -> list[dict[str, str]]:
    """Pull Key/Value pairs from row cells; supports multi-label rows."""
    pairs: list[dict[str, str]] = []
    pending_key = ""
    for row in rows:
        cells = [c.strip() for c in row if str(c or "").strip()]
        if not cells:
            continue
        # Flatten "Document No.: | GL-… | Version No.: | 1.0"
        i = 0
        while i < len(cells):
            cell = cells[i]
            if ":" in cell:
                key = cell.split(":", 1)[0].strip().rstrip(".")
                inline = cell.split(":", 1)[1].strip()
                if inline and is_valid_meta_key(key):
                    # May also contain next label: "GL-CQA-GOP-0030 Version No.:"
                    # Leave inline as value; extra label tokens handled by colon parse later.
                    pairs.append({"key": key, "value": inline})
                    pending_key = ""
                    i += 1
                    continue
                if is_valid_meta_key(key):
                    if i + 1 < len(cells):
                        pairs.append({"key": key, "value": cells[i + 1]})
                        pending_key = ""
                        i += 2
                        continue
                    pending_key = key
                    i += 1
                    continue
            if pending_key:
                pairs.append({"key": pending_key, "value": cell})
                pending_key = ""
                i += 1
                continue
            if i + 1 < len(cells) and is_valid_meta_key(cell.rstrip(":")):
                pairs.append({"key": cell.rstrip(":"), "value": cells[i + 1]})
                i += 2
                continue
            i += 1
        # Recover embedded "ID Version No.:" patterns via colon field parser.
        blob = " ".join(cells)
        for item in colon_fields_from_page_lines([blob]):
            name = str(item.get("name") or "").strip()
            value = str(item.get("value") or "").strip()
            if name and value:
                pairs.append({"key": name, "value": value})

    # Dedupe preserving order.
    out: list[dict[str, str]] = []
    seen: set[str] = set()
    for pair in pairs:
        key = str(pair.get("key") or "").strip()
        value = str(pair.get("value") or "").strip()
        if not key or not value or not is_valid_meta_key(key):
            continue
        mark = f"{key.casefold()}::{value}"
        if mark in seen:
            continue
        seen.add(mark)
        out.append({"key": key, "value": value})
    return out


_SIGNATURE_HEADING_RE = re.compile(
    r"\b("
    r"signatures?"
    r"|approvals?"
    r"|authorization"
    r"|authorisation"
    r"|document\s+approval"
    r"|prepared\s*/?\s*reviewed\s*/?\s*approved"
    r")\b",
    re.IGNORECASE,
)


def _has_signature_heading(table: dict[str, Any]) -> bool:
    """True only when the page-1 table carries an explicit signature/approval heading."""
    title = " ".join(str(table.get("title") or "").split()).strip()
    if title and _SIGNATURE_HEADING_RE.search(title):
        return True
    return False


def _pick_signature_table(tables: list[dict[str, Any]]) -> dict[str, Any]:
    """
    Choose the signature grid from page-1 tables only.

    Requires an explicit signature/approval heading (e.g. ``Signatures``).
    If that heading is missing, return empty — do not leak Abbreviations /
    classification / procedure grids into ``signature_table``.
    """
    empty = {"title": "", "columns": [], "rows": []}
    grids = [
        (idx, t)
        for idx, t in enumerate(tables)
        if str(t.get("kind") or "") == "grid"
        and (t.get("columns") or [])
        and (t.get("rows") or [])
        and _has_signature_heading(t)
    ]
    if not grids:
        return empty

    def score(item: tuple[int, dict[str, Any]]) -> tuple[int, int, int]:
        idx, table = item
        cols = len(table.get("columns") or [])
        rows = len(table.get("rows") or [])
        title = str(table.get("title") or "").strip()
        pts = 0
        # Compact person/approval tables usually have few columns and few rows.
        if 3 <= cols <= 6:
            pts += 20
        if 2 <= rows <= 12:
            pts += 20
        if rows > 20:
            pts -= 40
        if idx <= 2:
            pts += 15 - idx * 3
        if title and len(title.split()) <= 3 and ":" not in title:
            pts += 8
        # Demote tables whose first data cell looks like an outline number (6.7.1.).
        first_row = (table.get("rows") or [{}])[0]
        first_val = next(iter(first_row.values()), "") if isinstance(first_row, dict) else ""
        if re.match(r"^\d+(\.\d+)+\.?\s*", str(first_val or "")):
            pts -= 25
        return (pts, -idx, rows)

    _idx, chosen = max(grids, key=score)
    return {
        "title": str(chosen.get("title") or ""),
        "columns": list(chosen.get("columns") or []),
        "rows": list(chosen.get("rows") or []),
    }


def _fields_from_table_cells(rows: list[list[dict[str, Any]]]) -> list[dict[str, str]]:
    """Pull Label|Value pairs from tables using ':' morphology when present."""
    from .page_meta import is_valid_meta_key

    fields: list[dict[str, str]] = []
    lines: list[str] = []
    for row in rows:
        cells = [" ".join(str(c.get("text") or "").split()).strip() for c in row]
        unique_cells: list[str] = []
        for cell in cells:
            if not cell:
                continue
            if unique_cells and unique_cells[-1] == cell:
                continue
            unique_cells.append(cell)
        if not unique_cells:
            continue
        if len(unique_cells) == 1:
            lines.append(unique_cells[0])
            continue
        # Wide rows (signature grids, matrices) — only keep explicit Key: Value cells.
        if len(unique_cells) > 2:
            for cell in unique_cells:
                if ":" in cell:
                    lines.append(cell)
            continue
        # Exactly two cells: Label | Value (or Key: … morphology).
        left, right = unique_cells[0], unique_cells[1]
        if ":" in left:
            label = left.split(":", 1)[0].strip().rstrip(".")
            inline_val = left.split(":", 1)[1].strip()
            value = inline_val or right.strip()
            if is_valid_meta_key(label) and value:
                fields.append({"label": label, "name": label, "value": value})
            continue
        label = left.rstrip(":").strip()
        value = right.strip()
        if (
            is_valid_meta_key(label)
            and value
            and not (is_valid_meta_key(value) and not any(ch.isdigit() for ch in value))
            and not re.fullmatch(r"page(\s+of)?", value, flags=re.I)
        ):
            fields.append({"label": label, "name": label, "value": value})
    fields.extend(_colon_fields_from_lines(lines))
    return _unique_fields(fields)


def _extract_tables(doc: Document) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    return _extract_tables_from_container(doc.tables, location="body", section_index=None)


def _extract_tables_from_container(
    tables,
    *,
    location: str,
    section_index: int | None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    extracted: list[dict[str, Any]] = []
    images: list[dict[str, Any]] = []
    for table_index, table in enumerate(tables):
        rows_data: list[list[dict[str, Any]]] = []
        max_cols = 0
        for row_index, row in enumerate(table.rows):
            cells_out: list[dict[str, Any]] = []
            for col_index, cell in enumerate(row.cells):
                text = " ".join((cell.text or "").split()).strip()
                cell_images = _images_from_paragraphs(
                    cell.paragraphs,
                    location=location if "table" in location else f"{location}_table",
                    section_index=section_index,
                    table_index=table_index,
                    row_index=row_index,
                    col_index=col_index,
                )
                images.extend(cell_images)
                cells_out.append(
                    {
                        "text": text,
                        "row": row_index,
                        "column": col_index,
                        "images": cell_images,
                    }
                )
            max_cols = max(max_cols, len(cells_out))
            rows_data.append(cells_out)
        extracted.append(
            {
                "index": table_index,
                "location": location,
                "section_index": section_index,
                "rows": len(rows_data),
                "columns": max_cols,
                "cells": rows_data,
            }
        )
    return extracted, images
