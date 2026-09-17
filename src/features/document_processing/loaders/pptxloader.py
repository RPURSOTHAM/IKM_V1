"""PowerPoint (.pptx) loader — one Block group per slide."""

from __future__ import annotations

from pathlib import Path
from typing import Union

from src.features.document_processing.loaders.loader import Block, DocumentLoader


class PptxLoader(DocumentLoader):
    """Extract visible text from PowerPoint slides via python-pptx."""

    def load(self, path: Union[str, Path]) -> list[Block]:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError as exc:
            raise RuntimeError(
                "PowerPoint support requires python-pptx. Install python-pptx."
            ) from exc

        document_path = Path(path)
        presentation = Presentation(str(document_path))
        blocks: list[Block] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            line_no = 0
            for shape in slide.shapes:
                texts = _shape_texts(shape, MSO_SHAPE_TYPE)
                for text in texts:
                    cleaned = text.strip()
                    if not cleaned:
                        continue
                    line_no += 1
                    blocks.append(
                        Block(
                            text=cleaned,
                            page=slide_index,
                            line_number=line_no,
                            block_type="text",
                            component_type="paragraph",
                            metadata={"slide": slide_index, "source": "pptx"},
                        )
                    )
        return blocks


def _shape_texts(shape, mso_shape_type) -> list[str]:
    texts: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            line = "".join(run.text or "" for run in paragraph.runs).strip()
            if not line:
                line = (paragraph.text or "").strip()
            if line:
                texts.append(line)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                texts.append(row_text)
    if shape.shape_type == mso_shape_type.GROUP:
        for child in shape.shapes:
            texts.extend(_shape_texts(child, mso_shape_type))
    return texts
