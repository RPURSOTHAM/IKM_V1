"""Direct DOCX Citation & Page Provenance Manager.

Extracts, finds, formats, and manages numeric citations and page references directly
within DOCX files using python-docx without PDF conversion or external XML transforms.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Sequence

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    Document = None  # type: ignore[assignment]
    qn = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]
    RGBColor = None  # type: ignore[assignment]
    Table = None  # type: ignore[assignment]
    Paragraph = None  # type: ignore[assignment]

# Regex Patterns for Citation Formats
NUMERIC_CITATION_REGEX = re.compile(r"\[(\d+(?:\s*[-–,]\s*\d+)*)\]")
PAGE_CITATION_REGEX = re.compile(
    r"\b(?:p\.|pp\.|page|pages)\s*(\d+(?:\s*[-–]\s*\d+)?)\b",
    re.IGNORECASE,
)
NUMBERS_AND_PAGES_REGEX = re.compile(
    r"\[\d+(?:\s*[-–,]\s*\d+)*\]|\b(?:p\.|pp\.|page|pages)\s*\d+(?:\s*[-–]\s*\d+)?\b",
    re.IGNORECASE,
)
APA_CITATION_REGEX = re.compile(
    r"\(([A-Z][a-zA-Z\s\.\-&]+(?:et\s*al\.)?,\s*\d{4}[a-z]?(?:\s*;\s*[A-Z][a-zA-Z\s\.\-&]+(?:et\s*al\.)?,\s*\d{4}[a-z]?)*)\)"
)


class DocxDirectCitationManager:
    """Manages and extracts citations and page references directly from a DOCX file."""

    def __init__(self, docx_path_or_doc: str | Path | Any | None = None) -> None:
        if Document is None:
            raise RuntimeError(
                "python-docx is required for DocxDirectCitationManager. Run: pip install python-docx"
            )
        if docx_path_or_doc is None:
            self.docx_path = None
            self.doc = Document()
        elif isinstance(docx_path_or_doc, (str, Path)):
            self.docx_path = Path(docx_path_or_doc)
            self.doc = Document(str(self.docx_path))
        else:
            self.docx_path = None
            self.doc = docx_path_or_doc

    # ---------------------------------------------------------
    # 1. Direct Page Tracking (Word OpenXML Native Markers)
    # ---------------------------------------------------------
    @staticmethod
    def _element_has_page_break(element: Any) -> bool:
        """Check for explicit or Word-rendered page breaks in an XML element."""
        if qn is None:
            return False

        if element.tag == qn("w:p"):
            pPr = element.find(qn("w:pPr"))
            if pPr is not None and pPr.find(qn("w:pageBreakBefore")) is not None:
                return True

            for r in element.findall(qn("w:r")):
                for br in r.findall(qn("w:br")):
                    if br.get(qn("w:type")) == "page":
                        return True
                if r.find(qn("w:lastRenderedPageBreak")) is not None:
                    return True
        return False

    def scan_paragraphs_with_page_numbers(self) -> list[dict[str, Any]]:
        """Iterate through the DOCX body elements and track native page numbers."""
        items: list[dict[str, Any]] = []
        current_page = 1
        para_idx = 0

        for element in self.doc.element.body:
            if qn is not None and element.tag == qn("w:p"):
                para = Paragraph(element, self.doc)
                if self._element_has_page_break(element) and items:
                    current_page += 1

                text = para.text.strip()
                if text:
                    items.append({
                        "type": "paragraph",
                        "index": para_idx,
                        "page": current_page,
                        "text": text,
                        "paragraph_obj": para,
                    })
                para_idx += 1

            elif qn is not None and element.tag == qn("w:tbl"):
                table = Table(element, self.doc)
                for r_idx, row in enumerate(table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        cell_text = cell.text.strip()
                        if cell_text:
                            items.append({
                                "type": "table_cell",
                                "page": current_page,
                                "text": cell_text,
                                "row": r_idx,
                                "col": c_idx,
                                "cell_obj": cell,
                            })

        return items

    # ---------------------------------------------------------
    # 2. Find Citations (By Number & Page References)
    # ---------------------------------------------------------
    def find_citations(
        self,
        pattern: re.Pattern[str] | str = NUMERIC_CITATION_REGEX,
    ) -> list[dict[str, Any]]:
        """Find citations matching a regex pattern with page provenance."""
        if isinstance(pattern, str):
            compiled_pattern = re.compile(pattern)
        else:
            compiled_pattern = pattern

        scanned = self.scan_paragraphs_with_page_numbers()
        extracted: list[dict[str, Any]] = []

        for item in scanned:
            text = item["text"]
            for match in compiled_pattern.finditer(text):
                raw = match.group(0)
                extracted_numbers = [int(n) for n in re.findall(r"\d+", raw)]
                extracted.append({
                    "citation": raw,
                    "numbers": extracted_numbers,
                    "page": item["page"],
                    "element_type": item["type"],
                    "span": match.span(),
                    "context_text": text,
                })

        return extracted

    # ---------------------------------------------------------
    # 3. Add Cited Paragraphs directly to DOCX
    # ---------------------------------------------------------
    def add_cited_paragraph(
        self,
        text: str,
        citation_number: int | None = None,
        page_no: int | None = None,
        superscript: bool = True,
        color_rgb: tuple[int, int, int] | None = (0, 51, 102),
    ) -> Any:
        """Append a paragraph with a formatted citation number and/or page reference."""
        p = self.doc.add_paragraph()
        p.add_run(text.rstrip() + " ")

        if citation_number is not None:
            cite_run = p.add_run(f"[{citation_number}]")
            cite_run.font.superscript = superscript
            cite_run.font.bold = True
            if color_rgb is not None and RGBColor is not None:
                cite_run.font.color.rgb = RGBColor(*color_rgb)

        if page_no is not None:
            page_run = p.add_run(f" (p. {page_no})")
            page_run.font.italic = True
            if RGBColor is not None:
                page_run.font.color.rgb = RGBColor(100, 100, 100)

        return p

    # ---------------------------------------------------------
    # 4. Append References & Page Index Table
    # ---------------------------------------------------------
    def append_reference_table(
        self,
        references: Sequence[dict[str, Any]],
        title: str = "References & Page Index",
    ) -> Any:
        """Append a formatted references table at the end of the document."""
        self.doc.add_page_break()
        self.doc.add_heading(title, level=2)

        table = self.doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"

        hdr = table.rows[0].cells
        hdr[0].text = "#"
        hdr[1].text = "Source Title / Evidence"
        hdr[2].text = "Author / Section"
        hdr[3].text = "Page"

        for ref in references:
            row = table.add_row().cells
            row[0].text = f"[{ref.get('num', '')}]" if ref.get("num") is not None else ""
            row[1].text = str(ref.get("title", ref.get("evidence", "")))
            row[2].text = str(ref.get("author", ref.get("section", "")))
            page_val = ref.get("page")
            row[3].text = f"p. {page_val}" if page_val is not None else ""

        return table

    def save(self, output_path: str | Path) -> None:
        """Save the DOCX file directly."""
        self.doc.save(str(output_path))


def find_citations_in_docx(
    file_name: str | Path,
    pattern: re.Pattern[str] | str = NUMERIC_CITATION_REGEX,
) -> list[dict[str, Any]]:
    """Standalone helper to find citations and page provenance in a DOCX file."""
    manager = DocxDirectCitationManager(file_name)
    return manager.find_citations(pattern)


def convert_apa_to_numbered_docx(
    input_docx: str | Path,
    output_docx: str | Path,
) -> dict[str, int]:
    """Convert APA-style author-year citations to bracketed numbers [1], [2] in DOCX."""
    manager = DocxDirectCitationManager(input_docx)
    citation_map: dict[str, int] = {}
    current_num = 1

    for para in manager.doc.paragraphs:
        matches = list(APA_CITATION_REGEX.finditer(para.text))
        if not matches:
            continue

        new_text = para.text
        for match in reversed(matches):
            apa_citation = match.group(1).strip()
            if apa_citation not in citation_map:
                citation_map[apa_citation] = current_num
                current_num += 1

            num_str = f"[{citation_map[apa_citation]}]"
            start, end = match.span()
            new_text = new_text[:start] + num_str + new_text[end:]

        para.text = new_text

    if citation_map:
        references = [
            {"num": num, "title": apa_text, "section": "Literature", "page": ""}
            for apa_text, num in sorted(citation_map.items(), key=lambda x: x[1])
        ]
        manager.append_reference_table(references, title="References")

    manager.save(output_docx)
    return citation_map


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2 or sys.argv[1] in ("-h", "--help"):
        print("Usage: python src/features/citations/application/docx_citation_manager.py <path_to_docx_file>")
        sys.exit(0)

    target_file = Path(sys.argv[1])
    if not target_file.exists():
        print(f"Error: File not found: {target_file}")
        sys.exit(1)

    print(f"\n========================================================")
    print(f" Scanning DOCX Citations: {target_file.name}")
    print(f"========================================================\n")

    # 1. Search for Numeric Citations [1], [2, 3]
    numeric_results = find_citations_in_docx(target_file, NUMERIC_CITATION_REGEX)
    print(f"[+] Found {len(numeric_results)} numeric citation(s):")
    for r in numeric_results:
        print(f"    - Page {r['page']}: Citation {r['citation']} (Numbers: {r['numbers']})")
        print(f"      Context: \"{r['context_text'][:100]}\"...\n")

    # 2. Search for Page-Number References (p. 45, Page 12)
    page_results = find_citations_in_docx(target_file, PAGE_CITATION_REGEX)
    print(f"[+] Found {len(page_results)} page-number reference(s):")
    for r in page_results:
        print(f"    - Page {r['page']}: Reference \"{r['citation']}\"")
        print(f"      Context: \"{r['context_text'][:100]}\"...\n")

    # 3. Search for APA-style citations (Author, Year)
    apa_results = find_citations_in_docx(target_file, APA_CITATION_REGEX)
    print(f"[+] Found {len(apa_results)} APA-style citation(s):")
    for r in apa_results:
        print(f"    - Page {r['page']}: Citation {r['citation']}")
        print(f"      Context: \"{r['context_text'][:100]}\"...\n")
